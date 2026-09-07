"""generar_diapositivas_entrega1.py

Fecha: 2026-09-06
Descripcion: genera la presentacion PowerPoint (16:9) de la Entrega 1 del
proyecto PLN (Corpus, Idea de Proyecto y Textometria) a partir de los
resultados verificados (resultados/metricas.json, informe PDF) y de las
figuras PNG existentes. La salida es
diapositivas/Entrega1_Diapositivas.pptx (14 diapositivas).

Diseno (decisiones documentadas):
- Slide size 16:9 (13.333 x 7.5 in), plantilla en blanco.
- Banda de titulo azul oscuro #1F3864 con texto blanco 28 pt bold.
- Texto de cuerpo #333333; bullets 13-16 pt; encabezados de tabla sombreados.
- Figuras insertadas con add_picture conservando la proporcion real del PNG
  (dimensiones leidas con PIL) y escaladas a ~55-65 % del alto del slide,
  centradas en su zona disponible.
- Notas del orador (2-4 frases) en todas las diapositivas.
- Numero de diapositiva pequeno en el pie (a partir de la diapositiva 2).

Estado del arte aplicado: ninguno adicional; el contenido reproduce
exactamente las cifras verificadas de la Entrega 1 (no se inventan valores).
"""

from __future__ import annotations

import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

# ---------------------------------------------------------------------------
# Paleta y parametros de diseno
# ---------------------------------------------------------------------------
DARK_BLUE = RGBColor(0x1F, 0x38, 0x64)
SECONDARY_BLUE = RGBColor(0x2F, 0x5D, 0x8A)
TEXT_COLOR = RGBColor(0x33, 0x33, 0x33)
GRAY_COLOR = RGBColor(0x59, 0x59, 0x59)
LIGHT_GRAY = RGBColor(0x8A, 0x8A, 0x8A)
LIGHT_BG = RGBColor(0xEE, 0xF3, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
NOTE_BG = RGBColor(0xF2, 0xF5, 0xFA)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
FONT = "Calibri"

BASE_DIR = Path(__file__).resolve().parents[1]
FIG_DIR = BASE_DIR / "figuras"
INFO_DIR = BASE_DIR / "informe"
OUT_DIR = BASE_DIR / "diapositivas"
OUT_FILE = OUT_DIR / "Entrega1_Diapositivas.pptx"

# Rutas de las figuras usadas en la presentacion
FIG_TOP20_CON = FIG_DIR / "fig_top20_con_stopwords.png"
FIG_TOP20_SIN = FIG_DIR / "fig_top20_sin_stopwords.png"
FIG_POR_ANIO = FIG_DIR / "fig_sentencias_por_anio.png"
FIG_NUBE_GLOBAL = FIG_DIR / "fig_nube_global.png"
FIG_NUBES_1992_2001 = INFO_DIR / "fig_nubes_por_anio_1992_2001.png"
FIG_NUBES_2002_2011 = INFO_DIR / "fig_nubes_por_anio_2002_2011.png"
FIG_NUBES_2012_2020 = INFO_DIR / "fig_nubes_por_anio_2012_2020.png"


def png_size(path: Path) -> tuple[int, int]:
    """Devuelve (ancho, alto) en pixeles de un PNG usando PIL."""
    with Image.open(path) as im:
        return im.size


# ---------------------------------------------------------------------------
# Utilidades de texto (estimacion de desborde y parrafos)
# ---------------------------------------------------------------------------
def est_lineas(texto: str, size_pt: float, width_in: float) -> int:
    """Estimacion conservadora de lineas que ocupara un texto al envolverse."""
    if not texto:
        return 0
    # factor de ancho de caracter para Calibri (~0.53 x tamano en puntos)
    ancho_char = 0.55 * size_pt / 72.0
    chars_por_linea = max(1, int(width_in / ancho_char))
    total = 0
    for segmento in texto.split("\n"):
        total += max(1, -(-len(segmento) // chars_por_linea))
    return total


def agregar_parrafo(tf, texto, size=15, bold=False, color=TEXT_COLOR,
                    space_after=6, space_before=0, align=PP_ALIGN.LEFT,
                    bullet=None, italic=False, first=False, line_spacing=1.0):
    """Agrega un parrafo con un run unico al text frame dado."""
    p = tf.paragraphs[0] if first and not tf.paragraphs[0].runs else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    p.space_before = Pt(space_before)
    if line_spacing:
        p.line_spacing = line_spacing
    run = p.add_run()
    texto_final = texto
    if bullet:
        texto_final = bullet + texto_final
    run.text = texto_final
    f = run.font
    f.name = FONT
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color
    return p


def agregar_parrafo_dual(tf, lead, resto, size=15, color=TEXT_COLOR,
                         space_after=6, bullet="\u2022  ", first=False):
    """Parrafo con encabezado en negrita (lead) y cuerpo normal (resto)."""
    p = tf.paragraphs[0] if first and not tf.paragraphs[0].runs else tf.add_paragraph()
    p.space_after = Pt(space_after)
    p.line_spacing = 1.04
    if bullet:
        run_b = p.add_run()
        run_b.text = bullet
        run_b.font.name = FONT
        run_b.font.size = Pt(size)
        run_b.font.color.rgb = color
    if lead:
        run_l = p.add_run()
        run_l.text = lead
        run_l.font.name = FONT
        run_l.font.size = Pt(size)
        run_l.font.bold = True
        run_l.font.color.rgb = color
    run_r = p.add_run()
    run_r.text = resto
    run_r.font.name = FONT
    run_r.font.size = Pt(size)
    run_r.font.color.rgb = color
    return p


# ---------------------------------------------------------------------------
# Utilidades de forma
# ---------------------------------------------------------------------------
def add_rect(slide, x, y, w, h, fill, line_color=None):
    """Rectangulo solido sin borde visible."""
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line_color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line_color
    shape.shadow.inherit = False
    return shape


def add_textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    """Caja de texto con word_wrap activo y anclaje configurable."""
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    return tf


def banda_titulo(slide, texto, size=28):
    """Banda azul oscuro superior con el titulo en blanco (bold)."""
    add_rect(slide, 0, 0, SLIDE_W, Inches(1.02), DARK_BLUE)
    box = slide.shapes.add_textbox(Inches(0.55), Inches(0.08),
                                   SLIDE_W - Inches(1.1), Inches(0.86))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = texto
    f = run.font
    f.name = FONT
    f.size = Pt(size)
    f.bold = True
    f.color.rgb = WHITE
    return box


def pie_diapositiva(slide, numero):
    """Pie con texto institucional y numero de diapositiva."""
    tf = add_textbox(slide, Inches(0.55), Inches(7.12), Inches(6.0), Inches(0.3))
    agregar_parrafo(tf, "PLN \u00b7 Entrega 1 \u00b7 Pontificia Universidad Javeriana",
                    size=9, color=LIGHT_GRAY, space_after=0, first=True)
    tf2 = add_textbox(slide, SLIDE_W - Inches(1.1), Inches(7.12), Inches(0.55), Inches(0.3))
    p = tf2.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    run = p.add_run()
    run.text = str(numero)
    run.font.name = FONT
    run.font.size = Pt(10)
    run.font.color.rgb = LIGHT_GRAY


def add_picture_scaled(slide, path, target_h_in, center_x_in, top_in):
    """Inserta un PNG con la altura pedida conservando su proporcion.

    Args:
        slide: Diapositiva destino.
        path: Ruta del PNG.
        target_h_in: Altura deseada en pulgadas.
        center_x_in: Centro horizontal en pulgadas.
        top_in: Borde superior en pulgadas.

    Returns:
        La forma picture creada.
    """
    w_px, h_px = png_size(path)
    w_in = target_h_in * (w_px / h_px)
    left = center_x_in - w_in / 2.0
    if not (0.0 <= left and left + w_in <= 13.333 and target_h_in > 0):
        raise ValueError(
            "Imagen fuera de rango: {} left={:.2f} right={:.2f} top={:.2f} h={:.2f}".format(
                path.name, left, left + w_in, top_in, target_h_in))
    pic = slide.shapes.add_picture(str(path), Inches(left), Inches(top_in),
                                   width=Inches(w_in), height=Inches(target_h_in))
    return pic


def notas_slide(slide, texto):
    """Establece las notas del orador de la diapositiva."""
    slide.notes_slide.notes_text_frame.text = texto


def verificar_ajuste(slide, indice, advertencias):
    """Advertencia (no bloqueante) si un text box parece no caber en su caja."""
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        tf = shape.text_frame
        texto = tf.text
        if not texto.strip():
            continue
        alto_in = shape.height / 914400.0
        ancho_in = shape.width / 914400.0
        alto_est = 0.0
        for p in tf.paragraphs:
            size_pt = max((r.font.size.pt if r.font.size else 14.0) for r in p.runs) \
                if p.runs else 14.0
            n_lineas = est_lineas(p.text, size_pt, ancho_in)
            spacing = p.line_spacing if p.line_spacing else 1.0
            alto_est += n_lineas * size_pt * 1.22 * spacing / 72.0
            alto_est += (p.space_after.pt if p.space_after else 0.0) / 72.0
        if alto_est > alto_in * 1.06:
            advertencias.append(
                "Slide {}: textbox probable desborde (est {} in > caja {} in): {}...".format(
                    indice, round(alto_est, 2), round(alto_in, 2), texto[:60].replace("\n", " ")))


# ---------------------------------------------------------------------------
# Construccion de las 14 diapositivas
# ---------------------------------------------------------------------------
def construir(prs: Presentation) -> list[str]:
    """Construye todas las diapositivas; devuelve advertencias de layout."""
    advertencias: list[str] = []
    blank = prs.slide_layouts[6]

    # =====================================================================
    # Diapositiva 1 - Portada
    # =====================================================================
    slide = prs.slides.add_slide(blank)
    notas_slide(slide, "Saludo: presentamos la Entrega 1 del proyecto de "
                       "Procesamiento de Lenguaje Natural. El proyecto analiza "
                       "computacionalmente un corpus publico de 23.742 sentencias "
                       "de la Corte Constitucional de Colombia (1992-2021). Esta "
                       "entrega deja el corpus depurado y una textometria sobre la "
                       "totalidad del corpus, base de las entregas siguientes.")

    add_rect(slide, 0, 0, SLIDE_W, Inches(0.28), SECONDARY_BLUE)
    add_rect(slide, 0, Inches(6.78), SLIDE_W, Inches(0.72), DARK_BLUE)

    tf = add_textbox(slide, Inches(0.8), Inches(0.55), SLIDE_W - Inches(1.6), Inches(0.4))
    agregar_parrafo(tf, "PONTIFICIA UNIVERSIDAD JAVERIANA",
                    size=14, color=SECONDARY_BLUE, align=PP_ALIGN.CENTER,
                    space_after=0, first=True)
    tf = add_textbox(slide, Inches(0.8), Inches(0.98), SLIDE_W - Inches(1.6), Inches(0.4))
    agregar_parrafo(tf, "Procesamiento de Lenguaje Natural \u00b7 Entrega 1",
                    size=15, color=GRAY_COLOR, align=PP_ALIGN.CENTER,
                    space_after=0, first=True)

    add_rect(slide, Inches(5.17), Inches(1.62), Inches(3.0), Inches(0.045), DARK_BLUE)

    tf = add_textbox(slide, Inches(0.8), Inches(1.9), SLIDE_W - Inches(1.6), Inches(1.5))
    agregar_parrafo(tf, "Corpus, idea de proyecto y textometr\u00eda",
                    size=36, bold=True, color=DARK_BLUE, align=PP_ALIGN.CENTER,
                    space_after=2, first=True)
    agregar_parrafo(tf, "An\u00e1lisis computacional de la jurisprudencia de la "
                        "Corte Constitucional de Colombia (1992\u20132021)",
                    size=17, color=TEXT_COLOR, align=PP_ALIGN.CENTER, space_after=0)

    tf = add_textbox(slide, Inches(0.8), Inches(3.5), SLIDE_W - Inches(1.6), Inches(0.4))
    agregar_parrafo(tf, "23.742 sentencias \u00b7 260.988.316 palabras \u00b7 "
                        "corpus p\u00fablico Hugging Face (CC-BY-4.0)",
                    size=14, italic=True, color=SECONDARY_BLUE,
                    align=PP_ALIGN.CENTER, space_after=0, first=True)

    tf = add_textbox(slide, Inches(0.8), Inches(4.05), SLIDE_W - Inches(1.6), Inches(2.4))
    agregar_parrafo(tf, "Integrantes", size=13, color=LIGHT_GRAY,
                    align=PP_ALIGN.CENTER, space_after=10, first=True)
    for nombre in ["Erika Valeria Rivera Jim\u00e9nez",
                   "Juli\u00e1n Guillermo Carrillo Meneses",
                   "Carlos Andr\u00e9s Garnica Salazar",
                   "Oscar Andr\u00e9s Prieto Moreno"]:
        agregar_parrafo(tf, nombre, size=18, color=TEXT_COLOR,
                        align=PP_ALIGN.CENTER, space_after=6)

    tf = add_textbox(slide, Inches(0.8), Inches(6.86), SLIDE_W - Inches(1.6), Inches(0.55))
    agregar_parrafo(tf, "Septiembre de 2026", size=16, color=WHITE,
                    align=PP_ALIGN.CENTER, space_after=0, first=True)
    verificar_ajuste(slide, 1, advertencias)

    # =====================================================================
    # Diapositiva 2 - Objetivo de la entrega
    # =====================================================================
    slide = prs.slides.add_slide(blank)
    banda_titulo(slide, "Objetivo de la entrega")
    notas_slide(slide, "Segun el enunciado, la Entrega 1 pide cuatro productos: la "
                       "idea de proyecto, el corpus descrito con evidencia, una "
                       "textometria basica y un informe de maximo 8 paginas con su "
                       "anexo. En las siguientes diapositivas mostramos como se "
                       "cumple cada punto, siempre con cifras verificables.")

    tf = add_textbox(slide, Inches(0.75), Inches(1.45), Inches(11.8), Inches(3.4))
    items_obj = [
        ("Idea de proyecto: ", "dominio (lenguaje jur\u00eddico / Corte Constitucional), "
                               "problema de PLN y preguntas de investigaci\u00f3n."),
        ("Corpus con evidencia: ", "fuente verificada (Hugging Face), licencia, formato, "
                                   "caracter\u00edsticas, depuraci\u00f3n y fragmentos "
                                   "representativos."),
        ("Textometr\u00eda b\u00e1sica: ", "pipeline de limpieza documentado, medidas textuales, "
                                         "top-20 con y sin stopwords, distribuci\u00f3n temporal "
                                         "y nubes de palabras, sobre el corpus completo."),
        ("Informe: ", "documento final de m\u00e1ximo 8 p\u00e1ginas (PDF) m\u00e1s anexo con "
                      "muestra del corpus (25 sentencias; semilla 42)."),
    ]
    for i, (lead, resto) in enumerate(items_obj):
        agregar_parrafo_dual(tf, lead, resto, size=17, space_after=14, first=(i == 0))

    add_rect(slide, Inches(0.75), Inches(5.15), Inches(11.8), Inches(1.35), NOTE_BG)
    tf = add_textbox(slide, Inches(1.05), Inches(5.38), Inches(11.2), Inches(0.95))
    agregar_parrafo(tf, "Entregables verificables", size=13, bold=True,
                    color=SECONDARY_BLUE, space_after=4, first=True)
    agregar_parrafo(tf, "Informe Entrega_1_Informe.pdf (8 p\u00e1ginas) y anexo CSV de muestra \u00b7 "
                        "notebook consolidado con el pipeline reproducible \u00b7 "
                        "figuras/ y resultados/metricas.json.",
                    size=13.5, color=TEXT_COLOR, space_after=0)
    pie_diapositiva(slide, 2)
    verificar_ajuste(slide, 2, advertencias)

    # =====================================================================
    # Diapositiva 3 - Idea de proyecto
    # =====================================================================
    slide = prs.slides.add_slide(blank)
    banda_titulo(slide, "Idea de proyecto: PLN sobre jurisprudencia constitucional")
    notas_slide(slide, "El proyecto trabaja sobre el lenguaje juridico: sentencias "
                       "de la Corte Constitucional de Colombia. El problema central es "
                       "extraer conocimiento de un acervo enorme sin metadatos "
                       "estructurados. Justificamos el enfoque por la inviabilidad de "
                       "la lectura manual y la escasez de recursos legales anotados "
                       "en espanol; cerramos con las 4 preguntas de investigacion.")

    tf = add_textbox(slide, Inches(0.75), Inches(1.32), Inches(11.8), Inches(3.15))
    bloques = [
        ("Dominio. ", "Lenguaje jur\u00eddico: jurisprudencia de la Corte Constitucional "
                      "colombiana; sentencias de tutela (T), constitucionalidad (C) y "
                      "unificaci\u00f3n (SU); documentos extensos (\u224811.000 palabras) "
                      "con estructura formal, f\u00f3rmulas de estilo estables y citas."),
        ("Problema. ", "Corpus extenso sin metadatos: 23.742 documentos y \u2248261 "
                       "millones de palabras, con solo identificador y texto; la "
                       "informaci\u00f3n relevante se dispersa entre f\u00f3rmulas "
                       "procesales repetidas y el tipo de providencia y el a\u00f1o "
                       "deben inferirse del propio texto."),
        ("Justificaci\u00f3n. ", "La lectura manual de decenas de miles de documentos "
                                 "largos es inviable; el PLN permite evidencia "
                                 "cuantificable y reproducible. Escasez documentada de "
                                 "recursos legales anotados en espa\u00f1ol "
                                 "(Latinoam\u00e9rica): un corpus depurado y "
                                 "caracterizado ya es un aporte."),
    ]
    for i, (lead, resto) in enumerate(bloques):
        agregar_parrafo_dual(tf, lead, resto, size=14.5, space_after=8, first=(i == 0))

    tf = add_textbox(slide, Inches(0.75), Inches(4.62), Inches(11.8), Inches(0.35))
    agregar_parrafo(tf, "Preguntas de investigaci\u00f3n",
                    size=14, bold=True, color=SECONDARY_BLUE, space_after=0, first=True)
    tf = add_textbox(slide, Inches(0.95), Inches(5.05), Inches(11.6), Inches(1.95))
    preguntas = [
        "\u00bfQu\u00e9 propiedades l\u00e9xicas y estilom\u00e9tricas caracterizan las "
        "sentencias y permiten diferenciar los tipos de providencia (T, C y SU) y los periodos?",
        "\u00bfC\u00f3mo evoluciona el vocabulario de la Corte entre 1992 y 2021 en riqueza "
        "l\u00e9xica, t\u00e9rminos dominantes y temas recurrentes?",
        "\u00bfEs posible clasificar autom\u00e1ticamente las sentencias seg\u00fan su tipo "
        "o tem\u00e1tica con t\u00e9cnicas de PLN?",
        "\u00bfEn qu\u00e9 medida los m\u00e9todos autom\u00e1ticos pueden apoyar la consulta "
        "de sentencias extensas mediante res\u00famenes extractivos o abstractivos?",
    ]
    for i, q in enumerate(preguntas):
        agregar_parrafo(tf, q, size=13.5, color=TEXT_COLOR, space_after=5,
                        bullet="\u2013  ", first=(i == 0))
    pie_diapositiva(slide, 3)
    verificar_ajuste(slide, 3, advertencias)

    # =====================================================================
    # Diapositiva 4 - Corpus: fuente, caracteristicas y depuracion
    # =====================================================================
    slide = prs.slides.add_slide(blank)
    banda_titulo(slide, "Corpus: fuente, caracter\u00edsticas y depuraci\u00f3n")
    notas_slide(slide, "El corpus se descargo del repositorio abierto de Hugging Face "
                       "Manuel/sentencias-corte-cons-colombia-1992-2021, licencia "
                       "CC-BY-4.0: un CSV de 1,68 GB con dos columnas. El nombre del "
                       "archivo sugiere 2022, pero la cobertura real llega a 2021, "
                       "parcial. De 23.750 filas quedaron 23.742 sentencias: el regex "
                       "de prefijo extrajo identificador y anio en el 99,97 %; las 8 "
                       "filas restantes (6 autos y 2 fragmentos HTML) se excluyeron y "
                       "documentaron.")

    tf = add_textbox(slide, Inches(0.75), Inches(1.32), Inches(7.6), Inches(5.6))
    items_corpus = [
        ("Fuente. ", "Repositorio p\u00fablico \u201csentencias-corte-cons-colombia-"
                     "1992-2021\u201d (usuario Manuel) en el Hugging Face Hub; licencia "
                     "CC-BY-4.0; textos de la relator\u00eda de la Corte Constitucional."),
        ("Formato. ", "Un \u00fanico CSV UTF-8 de 1,68 GB (Juris1992_2022.csv), con dos "
                      "columnas: identificador y texto completo; sin metadatos "
                      "estructurados."),
        ("Cobertura real. ", "1992\u20132021 (30 a\u00f1os): el archivo no contiene "
                             "sentencias de 2022 y 2021 es parcial (solo 20 documentos)."),
        ("Depuraci\u00f3n. ", "Expresi\u00f3n regular de prefijo para extraer "
                             "identificador y a\u00f1o: \u00e9xito en 23.742 de 23.750 "
                             "filas (99,97 %)."),
        ("Exclusiones. ", "8 filas excluidas y documentadas: 6 autos (no sentencias) y "
                          "2 fragmentos HTML residuales de la recolecci\u00f3n."),
        ("Fragmentos. ", "Ejemplos verificados en el anexo: T-293/97, C-448/98 y "
                         "SU-476/97."),
    ]
    for i, (lead, resto) in enumerate(items_corpus):
        agregar_parrafo_dual(tf, lead, resto, size=13.5, space_after=9, first=(i == 0))

    panel_x = Inches(8.7)
    add_rect(slide, panel_x, Inches(1.35), Inches(3.9), Inches(5.35), LIGHT_BG)
    add_rect(slide, panel_x, Inches(1.35), Inches(3.9), Inches(0.55), DARK_BLUE)
    tf = add_textbox(slide, panel_x + Inches(0.2), Inches(1.4), Inches(3.5), Inches(0.45),
                     anchor=MSO_ANCHOR.MIDDLE)
    agregar_parrafo(tf, "Corpus de an\u00e1lisis", size=15, bold=True, color=WHITE,
                    space_after=0, first=True)
    tf = add_textbox(slide, panel_x + Inches(0.25), Inches(2.15), Inches(3.45), Inches(4.4))
    filas_panel = [
        ("23.742", "sentencias depuradas"),
        ("17.099 T \u00b7 6.246 C \u00b7 397 SU", "tutela \u00b7 constitucionalidad \u00b7 unificaci\u00f3n"),
        ("1992\u20132021", "espa\u00f1ol \u00b7 registro jur\u00eddico"),
        ("\u224811.000 palabras", "por documento (promedio)"),
    ]
    for i, (valor, etiqueta) in enumerate(filas_panel):
        agregar_parrafo(tf, valor, size=17, bold=True, color=DARK_BLUE,
                        space_after=1, first=(i == 0))
        agregar_parrafo(tf, etiqueta, size=12, color=GRAY_COLOR, space_after=12)
    pie_diapositiva(slide, 4)
    verificar_ajuste(slide, 4, advertencias)

    # =====================================================================
    # Diapositiva 5 - Metodologia: pipeline de limpieza
    # =====================================================================
    slide = prs.slides.add_slide(blank)
    banda_titulo(slide, "Metodolog\u00eda: limpieza y textometr\u00eda")
    notas_slide(slide, "Decision clave: analizar el corpus completo (23.742 "
                       "sentencias) y no una muestra; la pasada completa tarda unos 5 "
                       "minutos en CPU y la semilla fija 42 se usa para la muestra del "
                       "anexo y las nubes. La limpieza tiene 8 pasos medibles y "
                       "documentados, inspirados en la evidencia de que el "
                       "preprocesamiento cambia los resultados; la lista de stopwords "
                       "se normaliza con el mismo criterio que el corpus.")

    tf = add_textbox(slide, Inches(0.75), Inches(1.25), Inches(11.8), Inches(0.6))
    agregar_parrafo_dual(tf, "Decisi\u00f3n metodol\u00f3gica. ",
                         "textometr\u00eda sobre el corpus completo (no muestral): "
                         "\u22485 min de procesamiento en CPU; semilla fija 42 para "
                         "la muestra del anexo y las nubes.",
                         size=14, space_after=0, bullet="", first=True)

    pasos_izq = [
        ("1", "Conversi\u00f3n del texto a min\u00fasculas."),
        ("2", "Protecci\u00f3n temporal de la \u00f1 para que no se pierda en el paso siguiente."),
        ("3", "Normalizaci\u00f3n Unicode (NFD) y eliminaci\u00f3n de diacr\u00edticos "
              "(\u00e1\u2192a, \u00e9\u2192e, etc.)."),
        ("4", "Filtro de caracteres: solo letras [a-z\u00f1]; se eliminan d\u00edgitos, "
              "puntuaci\u00f3n y s\u00edmbolos."),
    ]
    pasos_der = [
        ("5", "Colapso de espacios m\u00faltiples."),
        ("6", "Tokenizaci\u00f3n por separaci\u00f3n en espacios."),
        ("7", "Eliminaci\u00f3n de tokens de una letra (restos de claves T-/C-/SU-) y de "
              "la abreviatura de cita \u201cart\u201d."),
        ("8", "Stopwords NLTK del espa\u00f1ol normalizadas con el mismo criterio: "
              "quedan 306 formas (84 con tilde corregidas)."),
    ]

    for col_x, pasos in ((0.95, pasos_izq), (6.95, pasos_der)):
        tf = add_textbox(slide, Inches(col_x), Inches(2.15), Inches(5.5), Inches(3.6))
        for i, (num, texto) in enumerate(pasos):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(12)
            p.line_spacing = 1.05
            rn = p.add_run()
            rn.text = num + ".  "
            rn.font.name = FONT
            rn.font.size = Pt(15)
            rn.font.bold = True
            rn.font.color.rgb = SECONDARY_BLUE
            rt = p.add_run()
            rt.text = texto
            rt.font.name = FONT
            rt.font.size = Pt(14.5)
            rt.font.color.rgb = TEXT_COLOR

    add_rect(slide, Inches(0.75), Inches(5.95), Inches(11.8), Inches(0.85), NOTE_BG)
    tf = add_textbox(slide, Inches(1.05), Inches(6.1), Inches(11.2), Inches(0.6))
    agregar_parrafo(tf, "La misma transformaci\u00f3n se aplica al corpus y a la lista de "
                        "stopwords: as\u00ed el filtro es consistente y no quedan formas "
                        "acentuadas sin eliminar (m\u00e1s\u2192mas, c\u00f3mo\u2192como).",
                    size=13, color=TEXT_COLOR, space_after=0, first=True)
    pie_diapositiva(slide, 5)
    verificar_ajuste(slide, 5, advertencias)

    # =====================================================================
    # Diapositiva 6 - Medidas textuales (tabla)
    # =====================================================================
    slide = prs.slides.add_slide(blank)
    banda_titulo(slide, "Medidas textuales del corpus")
    notas_slide(slide, "La tabla resume las medidas calculadas sobre las 23.742 "
                       "sentencias: 260.988.316 palabras crudas y 250.012.101 limpias, "
                       "con promedios de alrededor de 10.500 a 11.000 palabras por "
                       "documento. Es importante senalar que un valor previo de "
                       "271.907.978 palabras se descarto porque era una extrapolacion "
                       "de una muestra de 1.000 documentos, no un conteo real.")

    filas = [
        ("Documentos", "23.742"),
        ("Palabras crudas (raw)", "260.988.316"),
        ("Palabras normalizadas (limpias)", "250.012.101"),
        ("Promedio de palabras crudas por documento", "10.992,68"),
        ("Promedio de palabras limpias por documento", "10.530,37"),
    ]
    tabla_x = Inches(2.42)
    tabla_w = Inches(8.5)
    shape_tabla = slide.shapes.add_table(
        6, 2, tabla_x, Inches(1.5), tabla_w, Inches(3.35))
    tabla = shape_tabla.table
    tabla.columns[0].width = Inches(5.7)
    tabla.columns[1].width = Inches(2.8)
    tabla.rows[0].height = Inches(0.6)
    for i in range(1, 6):
        tabla.rows[i].height = Inches(0.55)

    def set_cell(celda, texto, size, bold, color, fill, align=PP_ALIGN.LEFT):
        celda.fill.solid()
        celda.fill.fore_color.rgb = fill
        celda.margin_left = Inches(0.14)
        celda.margin_right = Inches(0.14)
        celda.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf_c = celda.text_frame
        p = tf_c.paragraphs[0]
        p.alignment = align
        run = p.add_run()
        run.text = texto
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color

    set_cell(tabla.cell(0, 0), "Medida", 16, True, WHITE, DARK_BLUE)
    set_cell(tabla.cell(0, 1), "Valor", 16, True, WHITE, DARK_BLUE,
             align=PP_ALIGN.RIGHT)
    for i, (nombre, valor) in enumerate(filas, start=1):
        fill = LIGHT_BG if i % 2 == 1 else WHITE
        set_cell(tabla.cell(i, 0), nombre, 15, False, TEXT_COLOR, fill)
        set_cell(tabla.cell(i, 1), valor, 15, True, DARK_BLUE, fill,
                 align=PP_ALIGN.RIGHT)

    tf = add_textbox(slide, Inches(0.75), Inches(5.15), Inches(11.8), Inches(1.8))
    notas_tabla = [
        ("\u00bfQu\u00e9 se cont\u00f3? ", "Conteos reales sobre las 23.742 sentencias "
                                        "(corpus completo); procesamiento \u22485 min en CPU local."),
        ("Valor descartado. ", "Un total previo de 271.907.978 palabras crudas se "
                               "descart\u00f3: era una extrapolaci\u00f3n de una muestra "
                               "de 1.000 documentos (\u22484 % mayor que el valor real)."),
        ("Contraste. ", "Tras retirar las stopwords quedan 128.806.664 tokens: las "
                        "palabras funcionales pesan \u224848 % del corpus limpio."),
    ]
    for i, (lead, resto) in enumerate(notas_tabla):
        agregar_parrafo_dual(tf, lead, resto, size=14, space_after=7, first=(i == 0))
    pie_diapositiva(slide, 6)
    verificar_ajuste(slide, 6, advertencias)

    # =====================================================================
    # Diapositiva 7 - Frecuencias con stopwords
    # =====================================================================
    slide = prs.slides.add_slide(blank)
    banda_titulo(slide, "Frecuencias: top-20 con stopwords")
    notas_slide(slide, "Con las palabras vacias el top-5 lo encabezan de (22,5 "
                       "millones), la, el, que y en: puro andamiaje gramatical. Las "
                       "palabras funcionales representan cerca del 48 % de los tokens "
                       "limpios y la caida brusca de las frecuencias es consistente "
                       "con la ley de Zipf.")

    tf = add_textbox(slide, Inches(0.7), Inches(1.4), Inches(5.55), Inches(5.4))
    top5 = [
        ("1. de ", "22.543.375"),
        ("2. la ", "14.299.738"),
        ("3. el ", "8.225.551"),
        ("4. que ", "7.505.218"),
        ("5. en ", "7.208.184"),
    ]
    agregar_parrafo(tf, "Top-5 del corpus completo", size=16, bold=True,
                    color=DARK_BLUE, space_after=8, first=True)
    for rank, (palabra, valor) in enumerate(top5):
        p = tf.add_paragraph()
        p.space_after = Pt(4)
        rn = p.add_run()
        rn.text = palabra
        rn.font.name = FONT
        rn.font.size = Pt(18)
        rn.font.bold = True
        rn.font.color.rgb = SECONDARY_BLUE
        rt = p.add_run()
        rt.text = valor
        rt.font.name = FONT
        rt.font.size = Pt(18)
        rt.font.color.rgb = TEXT_COLOR

    items_zipf = [
        ("Peso funcional. ", "las stopwords son \u224848 % de los tokens limpios "
                             "(250.012.101 \u2192 128.806.664)."),
        ("Ley de Zipf. ", "ca\u00edda brusca de frecuencias: unas pocas palabras "
                          "concentran la mayor parte de las ocurrencias."),
        ("Interpretaci\u00f3n. ", "el top-20 con stopwords no aporta contenido: por "
                                  "eso se analiza el corpus sin ellas."),
    ]
    for i, (lead, resto) in enumerate(items_zipf):
        agregar_parrafo_dual(tf, lead, resto, size=13.5, space_after=8, first=False)

    # Figura derecha (ancho libre ~6.3 in; alto 4.4 in => 59 % del alto del slide)
    add_picture_scaled(slide, FIG_TOP20_CON, 4.4, 9.9, 1.75)
    pie_diapositiva(slide, 7)
    verificar_ajuste(slide, 7, advertencias)

    # =====================================================================
    # Diapositiva 8 - Frecuencias sin stopwords
    # =====================================================================
    slide = prs.slides.add_slide(blank)
    banda_titulo(slide, "Frecuencias: top-20 sin stopwords")
    notas_slide(slide, "Sin las palabras funcionales emerge el lexico juridico "
                       "dominante: articulo y ley (citas normativas), derecho(s), "
                       "corte, sentencia, constitucional y tutela (el amparo domina "
                       "el tipo T). Las grafias sin tilde (articulo, accion) son "
                       "efecto de la normalizacion; los tokens de una letra se "
                       "filtraron en el pipeline.")

    tf = add_textbox(slide, Inches(0.7), Inches(1.4), Inches(5.55), Inches(5.4))
    agregar_parrafo(tf, "L\u00e9xico jur\u00eddico dominante", size=16, bold=True,
                    color=DARK_BLUE, space_after=8, first=True)
    items_lex = [
        ("articulo ", "947.616  (citas normativas)"),
        ("derecho ", "898.926"),
        ("ley ", "809.313  (citas normativas)"),
        ("corte ", "760.197"),
        ("sentencia ", "736.326"),
        ("tutela ", "649.302  (mecanismo de amparo, tipo T)"),
    ]
    for palabra, resto in items_lex:
        p = tf.add_paragraph()
        p.space_after = Pt(4)
        rn = p.add_run()
        rn.text = palabra
        rn.font.name = FONT
        rn.font.size = Pt(15)
        rn.font.bold = True
        rn.font.color.rgb = SECONDARY_BLUE
        rt = p.add_run()
        rt.text = resto
        rt.font.name = FONT
        rt.font.size = Pt(14)
        rt.font.color.rgb = TEXT_COLOR

    items_sin = [
        ("Normalizaci\u00f3n. ", "articulo, accion, constitucion y proteccion sin "
                                 "tilde por la normalizaci\u00f3n (paso 3 del pipeline)."),
        ("Filtro previo. ", "tokens de una letra (t, c, p\u2026) y \u201cart\u201d "
                            "eliminados en origen."),
        ("Contenido. ", "el corpus muestra el registro esperado de jurisprudencia "
                        "constitucional colombiana."),
    ]
    for i, (lead, resto) in enumerate(items_sin):
        agregar_parrafo_dual(tf, lead, resto, size=13.5, space_after=8,
                             first=False)

    add_picture_scaled(slide, FIG_TOP20_SIN, 4.4, 9.9, 1.75)
    pie_diapositiva(slide, 8)
    verificar_ajuste(slide, 8, advertencias)

    # =====================================================================
    # Diapositiva 9 - Distribucion temporal
    # =====================================================================
    slide = prs.slides.add_slide(blank)
    banda_titulo(slide, "Distribuci\u00f3n temporal (1992\u20132021)")
    notas_slide(slide, "El volumen es estable: entre 1999 y 2014 se emitieron unas "
                       "950-990 providencias por anio. Luego desciende de forma "
                       "progresiva: 763 en 2015 hasta 505 en 2020, y colapsa a 20 "
                       "documentos en 2021, evidencia de que la cobertura de ese anio "
                       "es parcial y de que el nombre del archivo (Juris1992_2022) no "
                       "corresponde con el contenido real. Por tipo: tutela 72,0 %, "
                       "constitucionalidad 26,3 % y unificacion 1,7 %.")

    tf = add_textbox(slide, Inches(0.75), Inches(1.22), Inches(11.8), Inches(1.2))
    agregar_parrafo_dual(tf, "Estabilidad 1999\u20132014: ", "entre 950 y 990 "
                         "providencias anuales.", size=14, space_after=4, first=True)
    agregar_parrafo_dual(tf, "Descenso 2015\u20132020: ", "de 763 (2015) a 505 (2020); "
                         "2021 colapsa a 20 documentos (cobertura parcial; no hay 2022).",
                         size=14, space_after=4)
    agregar_parrafo_dual(tf, "Composici\u00f3n por tipo: ", "tutela 72,0 % \u00b7 "
                         "constitucionalidad 26,3 % \u00b7 unificaci\u00f3n 1,7 %.",
                         size=14, space_after=0)

    # Figura ancha (ratio 2.65): alto 4.2 in => 56 % del alto del slide
    add_picture_scaled(slide, FIG_POR_ANIO, 4.2, 6.6665, 2.55)
    pie_diapositiva(slide, 9)
    verificar_ajuste(slide, 9, advertencias)

    # =====================================================================
    # Diapositiva 10 - Nube global
    # =====================================================================
    slide = prs.slides.add_slide(blank)
    banda_titulo(slide, "Nube global de palabras")
    notas_slide(slide, "La nube se construye con los 5.000 terminos mas frecuentes "
                       "del corpus completo, sin stopwords gramaticales (306 formas "
                       "NLTK normalizadas) ni de dominio juridico-procesal (65 formas "
                       "de alta frecuencia como sentencia, corte, tutela o articulo). "
                       "Al retirar tambien el vocabulario administrativo del genero "
                       "quedan visibles los temas de fondo de la jurisprudencia.")

    tf = add_textbox(slide, Inches(0.75), Inches(1.22), Inches(11.8), Inches(1.0))
    items_nube = [
        ("Frecuencias. ", "se alimenta del top-5.000 de t\u00e9rminos m\u00e1s "
                          "frecuentes del corpus completo."),
        ("Listas retiradas. ", "stopwords gramaticales (306 formas NLTK "
                               "normalizadas) y 65 formas de dominio "
                               "jur\u00eddico-procesal de alta frecuencia."),
    ]
    for i, (lead, resto) in enumerate(items_nube):
        agregar_parrafo_dual(tf, lead, resto, size=13.5, space_after=5, first=(i == 0))

    # Figura ratio 1.71: alto 4.2 in => 56 % del alto del slide
    add_picture_scaled(slide, FIG_NUBE_GLOBAL, 4.2, 6.6665, 2.42)
    pie_diapositiva(slide, 10)
    verificar_ajuste(slide, 10, advertencias)

    # =====================================================================
    # Diapositiva 11 - Nubes anuales (3 paneles)
    # =====================================================================
    slide = prs.slides.add_slide(blank)
    banda_titulo(slide, "Nubes de palabras por a\u00f1o (3 segmentos)")
    notas_slide(slide, "Cada panel agrupa una nube por anio con los 1.000 terminos "
                       "mas frecuentes de ese anio (2021 se excluye por tener menos "
                       "de 50 documentos). Se presentan en tres paneles, 1992-2001, "
                       "2002-2011 y 2012-2020, para conservar legibilidad. La lectura "
                       "cualitativa muestra la estabilidad de las formulas del genero "
                       "juridico; el analisis cuantitativo de la evolucion lexica por "
                       "periodos se aborda en la entrega 2.")

    tf = add_textbox(slide, Inches(0.75), Inches(1.22), Inches(11.8), Inches(0.7))
    agregar_parrafo(tf, "Una nube por a\u00f1o con el top-1.000 de t\u00e9rminos de ese "
                        "a\u00f1o; 2021 se excluye por tener menos de 50 documentos "
                        "(criterio de agregaci\u00f3n).",
                    size=13.5, color=TEXT_COLOR, space_after=0, first=True)

    paneles = [
        (FIG_NUBES_1992_2001, "1992\u20132001"),
        (FIG_NUBES_2002_2011, "2002\u20132011"),
        (FIG_NUBES_2012_2020, "2012\u20132020"),
    ]
    # Imagenes muy altas (relacion w/h ~0.72): fila de 3 con alto 4.2 in cada una
    alto_panel = 4.2
    ancho_bruto = 0.0
    anchos = []
    for ruta, _ in paneles:
        w_px, h_px = png_size(ruta)
        w = alto_panel * (w_px / h_px)
        anchos.append(w)
        ancho_bruto += w
    espacio_gaps = Inches(0.35) * 2
    x_ini = (SLIDE_W - Inches(ancho_bruto) - espacio_gaps) / 2.0
    x_act = x_ini
    y_img = Inches(2.05)
    for (ruta, etiqueta), w in zip(paneles, anchos):
        pic = slide.shapes.add_picture(str(ruta), Emu(x_act), y_img,
                                       width=Inches(w), height=Inches(alto_panel))
        _ = pic  # ancho variable por panel
        x_act = Emu(x_act) + Emu(Inches(w)) + Inches(0.35)
        x_act = int(x_act)
    # Etiquetas debajo de cada panel
    x_act = x_ini
    for (ruta, etiqueta), w in zip(paneles, anchos):
        tf_lab = add_textbox(slide, Emu(int(x_act) + int((Inches(w) - Inches(2.2)) / 2)),
                             Inches(6.32), Inches(2.2), Inches(0.4))
        agregar_parrafo(tf_lab, etiqueta, size=14, bold=True, color=SECONDARY_BLUE,
                        align=PP_ALIGN.CENTER, space_after=0, first=True)
        x_act = int(x_act) + int(Inches(w)) + int(Inches(0.35))

    tf = add_textbox(slide, Inches(0.75), Inches(6.72), Inches(11.8), Inches(0.42))
    agregar_parrafo(tf, "Estabilidad cualitativa de las f\u00f3rmulas del g\u00e9nero; "
                        "el an\u00e1lisis diacr\u00f3nico se har\u00e1 en la entrega 2.",
                    size=12.5, italic=True, color=GRAY_COLOR,
                    align=PP_ALIGN.CENTER, space_after=0, first=True)
    pie_diapositiva(slide, 11)
    verificar_ajuste(slide, 11, advertencias)

    # =====================================================================
    # Diapositiva 12 - Estado del arte
    # =====================================================================
    slide = prs.slides.add_slide(blank)
    banda_titulo(slide, "Estado del arte: antecedentes")
    notas_slide(slide, "La busqueda (arXiv, Semantic Scholar, OpenAlex, Crossref y el "
                       "propio Hugging Face, verificada el 2026-09-05) no encontro "
                       "trabajos indexados que usen este dataset exacto. Los "
                       "antecedentes mas cercanos trabajan sobre el mismo corpus "
                       "subyacente de la relatoria o sobre metodologia de PLN legal "
                       "en espanol: Anzola Cortes, Barbosa Becerra, HaCohen-Kerner y "
                       "colaboradores, Gonzalez-Gonzalez y colaboradores, y "
                       "Chalkidis y colaboradores con LEGAL-BERT.")

    tf = add_textbox(slide, Inches(0.75), Inches(1.3), Inches(11.8), Inches(4.1))
    refs = [
        ("Anzola Cort\u00e9s (2021). ", "miner\u00eda de datos sobre m\u00e1s de 6.000 "
                                       "sentencias de constitucionalidad de la CCC "
                                       "(1992\u20132020): el mismo universo documental "
                                       "del dataset."),
        ("Barbosa Becerra (2021). ", "an\u00e1lisis l\u00e9xico-discursivo de sentencias "
                                     "de la CCC; evidencia la estabilidad de sus "
                                     "f\u00f3rmulas y un l\u00e9xico jur\u00eddico "
                                     "especializado."),
        ("HaCohen-Kerner et al. (2020). ", "la elecci\u00f3n del preprocesamiento cambia "
                                           "los resultados: fundamenta documentar cada "
                                           "paso de limpieza y normalizar corpus y "
                                           "stopwords por igual."),
        ("Gonz\u00e1lez-Gonz\u00e1lez et al. (2023). ", "clasificaci\u00f3n de sentencias "
                                                       "judiciales en espa\u00f1ol con "
                                                       "diccionarios de dominio "
                                                       "(\u201cexpert-in-the-loop\u201d)."),
        ("Chalkidis et al. (2020). ", "LEGAL-BERT: el vocabulario jur\u00eddico difiere "
                                      "del general; motiva evaluar adaptaci\u00f3n de "
                                      "dominio en las fases siguientes."),
    ]
    for i, (lead, resto) in enumerate(refs):
        agregar_parrafo_dual(tf, lead, resto, size=14, space_after=10, first=(i == 0))

    add_rect(slide, Inches(0.75), Inches(5.6), Inches(11.8), Inches(1.15), NOTE_BG)
    tf = add_textbox(slide, Inches(1.05), Inches(5.78), Inches(11.2), Inches(0.85))
    agregar_parrafo_dual(tf, "Nota. ", "no hay trabajos indexados (revistas, actas o "
                         "arXiv) que usen el dataset exacto de Hugging Face "
                         "(verificado 2026-09-05); los antecedentes cubren el mismo "
                         "corpus de la relator\u00eda y el PLN legal en espa\u00f1ol.",
                         size=13, space_after=0, bullet="", first=True)
    pie_diapositiva(slide, 12)
    verificar_ajuste(slide, 12, advertencias)

    # =====================================================================
    # Diapositiva 13 - Conclusiones
    # =====================================================================
    slide = prs.slides.add_slide(blank)
    banda_titulo(slide, "Conclusiones")
    notas_slide(slide, "Cumplimos el enunciado: idea de proyecto, corpus con "
                       "evidencia, textometria sobre el corpus completo e informe de "
                       "8 paginas con anexo. Las cifras clave confirman el caracter "
                       "extenso y formulaico del genero y sientan la linea base "
                       "medible para las siguientes entregas.")

    tf = add_textbox(slide, Inches(0.75), Inches(1.35), Inches(11.8), Inches(5.5))
    items_conc = [
        ("Cumplimiento. ", "idea de proyecto; corpus con evidencia (fuente, licencia, "
                           "formato, depuraci\u00f3n, fragmentos); textometr\u00eda "
                           "b\u00e1sica; informe de m\u00e1ximo 8 p\u00e1ginas con "
                           "anexo (25 sentencias, semilla 42)."),
        ("Corpus. ", "23.742 sentencias depuradas (T 17.099 \u00b7 C 6.246 \u00b7 "
                     "SU 397): extracci\u00f3n de identificador y a\u00f1o con regex en "
                     "el 99,97 % y 8 exclusiones documentadas (6 autos y 2 HTML)."),
        ("Textometr\u00eda real. ", "260.988.316 palabras crudas (\u224810.993 por "
                                   "documento) calculadas sobre el corpus completo; "
                                   "las stopwords pesan \u224848 % de los tokens limpios."),
        ("L\u00e9xico. ", "tras retirar las palabras funcionales domina el registro "
                          "jur\u00eddico-procesal: articulo, derecho, ley, corte, "
                          "sentencia, constitucional y tutela."),
        ("Distribuci\u00f3n temporal. ", "estable entre 1999 y 2014 (\u2248950\u2013990 "
                                        "por a\u00f1o), descenso 2015\u20132020 y 2021 "
                                        "parcial (20 documentos)."),
        ("Resultado. ", "el corpus es extenso y formulaico, como predice la literatura "
                        "sobre sentencias largas; esta textometr\u00eda es la l\u00ednea "
                        "base medible de la entrega 2."),
    ]
    for i, (lead, resto) in enumerate(items_conc):
        agregar_parrafo_dual(tf, lead, resto, size=14.5, space_after=11, first=(i == 0))
    pie_diapositiva(slide, 13)
    verificar_ajuste(slide, 13, advertencias)

    # =====================================================================
    # Diapositiva 14 - Trabajo futuro
    # =====================================================================
    slide = prs.slides.add_slide(blank)
    banda_titulo(slide, "Trabajo futuro (Entrega 2)")
    notas_slide(slide, "La entrega 2 avanzara en cuatro frentes: lematizacion y POS "
                       "con spaCy para agrupar flexiones; representaciones TF-IDF y "
                       "densas con medidas de diversidad lexica por anio y tipo; "
                       "clasificacion automatica por tipo y tematica comparando "
                       "modelos clasicos con LEGAL-BERT; y evaluacion de resumen "
                       "automatico para sentencias extensas.")

    tf = add_textbox(slide, Inches(0.75), Inches(1.5), Inches(11.8), Inches(4.6))
    items_fut = [
        ("Lematizaci\u00f3n y POS. ", "lematizaci\u00f3n y etiquetado de categor\u00edas "
                                      "gramaticales con spaCy para agrupar flexiones."),
        ("Representaciones. ", "TF-IDF y vectores densos (embeddings) con an\u00e1lisis "
                               "de diversidad l\u00e9xica por a\u00f1o y tipo de "
                               "providencia."),
        ("Clasificaci\u00f3n. ", "experimentos de clasificaci\u00f3n autom\u00e1tica por "
                                 "tipo (T/C/SU) y por tem\u00e1tica; contraste de modelos "
                                 "cl\u00e1sicos con modelos de dominio legal "
                                 "(p. ej., LEGAL-BERT)."),
        ("Res\u00famenes. ", "evaluaci\u00f3n de estrategias de resumen autom\u00e1tico "
                            "(extractivo o abstractivo) para apoyar la consulta de "
                            "sentencias extensas."),
    ]
    for i, (lead, resto) in enumerate(items_fut):
        agregar_parrafo_dual(tf, lead, resto, size=16, space_after=16, first=(i == 0))

    tf = add_textbox(slide, Inches(0.75), Inches(6.15), Inches(11.8), Inches(0.6))
    agregar_parrafo(tf, "La textometr\u00eda de esta entrega, con cifras sobre el corpus "
                        "completo y decisiones metodol\u00f3gicas documentadas, es la "
                        "l\u00ednea base de esas etapas.",
                    size=13.5, italic=True, color=GRAY_COLOR, space_after=0, first=True)
    pie_diapositiva(slide, 14)
    verificar_ajuste(slide, 14, advertencias)

    return advertencias


def main() -> None:
    """Genera el pptx, lo guarda y resume el resultado en consola."""
    try:
        prs = Presentation()
        prs.slide_width = SLIDE_W
        prs.slide_height = SLIDE_H
        prs.core_properties.title = "Entrega 1 - Corpus, idea de proyecto y textometria"
        prs.core_properties.subject = "Procesamiento de Lenguaje Natural - PUJ"
        prs.core_properties.author = "Rivera, Carrillo, Garnica y Prieto"

        advertencias = construir(prs)

        OUT_DIR.mkdir(parents=True, exist_ok=True)
        prs.save(str(OUT_FILE))

        print("Presentacion guardada en: {}".format(OUT_FILE))
        print("Diapositivas: {}".format(len(prs.slides._sldIdLst)))
        n_pics = 0
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.shape_type == 13:  # PICTURE
                    n_pics += 1
        print("Imagenes embebidas: {}".format(n_pics))
        print("Tamano: {:.2f} MB".format(OUT_FILE.stat().st_size / 1e6))
        if advertencias:
            print("ADVERTENCIAS DE LAYOUT (revisar):")
            for a in advertencias:
                print("  - " + a)
        else:
            print("Sin advertencias de desborde de texto estimadas.")
    except Exception as error:  # noqa: BLE001
        import traceback
        traceback.print_exc()
        sys.exit(1)


if __name__ == "__main__":
    main()
