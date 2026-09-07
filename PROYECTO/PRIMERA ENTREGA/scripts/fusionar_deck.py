"""fusionar_deck.py
Fecha: 2026-09-06
Descripción: Fusiona contenido (datos) del deck propio (diapositivas/Entrega1_Diapositivas.pptx)
en el deck del companero (PLN_PROYECTO.pptx), EDITANDO el deck del companero EN SU LUGAR y
conservando exactamente su plantilla/estilo/formato. No modifica ningun otro archivo del
proyecto; el deck propio solo se usa como referencia de contenido.

Cambios:
  A) Nueva slide de contenido "Contextualización / ¿Por qué es relevante?" (tras la slide
     "Objetivo y pregunta de investigación").
  B) Nueva slide "Construcción del corpus / Fragmentos representativos" con 3 tarjetas de
     citas literales (tras "Depuración y enriquecimiento del corpus").
  C) Nueva slide "Textometría / Nubes de palabras por año" con las 3 figuras de nubes anuales
     en fila (tras "Más allá del vocabulario jurídico").
  D) Nueva SECCION "Antecedentes": header de seccion + slide de estado del arte (5
     referencias + nota), antes del header "Conclusiones".
  E) Slide de procesamiento textual (paso 09): detalle ampliado de eliminacion de stopwords.
  F) Slide del Corpus (ficha técnica): dos lineas nuevas al final de la ficha.
  G) Slide 2 "Contenido" (agenda): 5 filas renumeradas/reordenadas con rangos reales.

Estado del arte aplicado: replicacion de estilos python-pptx 1.0.2 mediante clonado de
elementos XML de formas existentes del deck destino (fuentes, tamanos, colores, geometrias,
viñetas), garantizando fidelidad con la plantilla del companero.
"""
import copy
import math
import sys
import traceback

from PIL import Image
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.oxml.ns import qn
from lxml import etree

# =========================================================================
# CONFIG
# =========================================================================
CONFIG = {
    "deck_companero": "/home/carlos/Git/PLN/PROYECTO/PRIMERA ENTREGA/PLN_PROYECTO.pptx",
    "nubes": {
        "1992-2001": "/home/carlos/Git/PLN/PROYECTO/PRIMERA ENTREGA/informe/fig_nubes_por_anio_1992_2001.png",
        "2002-2011": "/home/carlos/Git/PLN/PROYECTO/PRIMERA ENTREGA/informe/fig_nubes_por_anio_2002_2011.png",
        "2012-2020": "/home/carlos/Git/PLN/PROYECTO/PRIMERA ENTREGA/informe/fig_nubes_por_anio_2012_2020.png",
    },
    "altura_nube_in": 4.05,
    "gap_nubes_in": 0.42,
}

EMU_IN = 914400.0
AZUL = "2F75B5"
GRIS_TEXTO = "425466"
GRIS_TARJETA = "162536"
PURPURA = "6554A4"
LILA = "EEEAF8"
AZUL_CLARO = "EAF2F8"


def _rgb(hexstr):
    from pptx.dml.color import RGBColor
    return RGBColor.from_string(hexstr)


# =========================================================================
# Helpers XML / slides
# =========================================================================
def max_shape_id(prs):
    mx = 0
    for s in prs.slides:
        for cnv in s.shapes._spTree.iter(qn("p:cNvPr")):
            try:
                mx = max(mx, int(cnv.get("id")))
            except (TypeError, ValueError):
                pass
    return mx


def clone_slide(prs, src_idx, shape_id_start):
    """Clona la slide src_idx (deepcopy del spTree + remapeo de relaciones de imagen)."""
    src = prs.slides[src_idx]
    dest = prs.slides.add_slide(src.slide_layout)
    for shp in list(dest.shapes):  # quitar placeholders instanciados por add_slide
        shp._element.getparent().remove(shp._element)
    for shp in src.shapes:
        dest.shapes._spTree.append(copy.deepcopy(shp._element))
    next_id = shape_id_start
    for cnv in dest.shapes._spTree.iter(qn("p:cNvPr")):
        cnv.set("id", str(next_id))
        next_id += 1
    for blip in dest.shapes._spTree.iter(qn("a:blip")):
        rid = blip.get(qn("r:embed"))
        if rid is None:
            continue
        rel = src.part.rels[rid]
        new_rid = dest.part.relate_to(rel.target_part, rel.reltype)
        blip.set(qn("r:embed"), new_rid)
    return dest


def find_shapes(slide, name=None, shape_type=None, min_w_in=0.0, has_text=None):
    out = []
    for sh in slide.shapes:
        if name is not None and sh.name != name:
            continue
        if shape_type is not None and sh.shape_type != shape_type:
            continue
        if sh.width is not None and sh.width / EMU_IN < min_w_in:
            continue
        if has_text is True and not (sh.has_text_frame and sh.text_frame.text.strip()):
            continue
        out.append(sh)
    return out


def remove_shape(slide, shp):
    shp._element.getparent().remove(shp._element)


def set_run_text(run, text):
    """Reemplaza el texto de un run conservando su rPr (formato)."""
    run.text = text


def first_content_run(paragraph):
    for r in paragraph.runs:
        if r.text.strip():
            return r
    return paragraph.runs[0] if paragraph.runs else None


def shape_paragraphs(shape):
    return shape.text_frame.paragraphs


def replace_band_title(slide, title):
    """Titulo sobre la banda: texto `title`, Arial 32pt bold blanco (patron del deck)."""
    for sh in slide.shapes:
        if sh.shape_type != 1 or not sh.has_text_frame:
            continue
        if not sh.name.startswith("Google Shape;105"):
            continue
        p = sh.text_frame.paragraphs[0]
        run = first_content_run(p)
        if run is None:
            return
        rpr = run._r.find(qn("a:rPr"))
        if rpr is not None:
            rpr.set("sz", "3200")
            rpr.set("b", "1")
        set_run_text(run, title)
        txBody = sh._element.find(qn("p:txBody"))
        bodypr = txBody.find(qn("a:bodyPr"))
        if bodypr is not None:
            for tag in ("a:normAutofit", "a:spAutoFit"):
                child = bodypr.find(qn(tag))
                if child is not None:
                    bodypr.remove(child)
            etree.SubElement(bodypr, qn("a:noAutofit"))
        return


def replace_subtitle(slide, subtitle):
    """Subtitulo (TextBox 7) conservando su formato 28pt bold ambar."""
    for sh in find_shapes(slide, name="TextBox 7", has_text=True):
        p = sh.text_frame.paragraphs[0]
        run = first_content_run(p)
        if run is not None:
            set_run_text(run, subtitle)
        return


# -------------------------------------------------------------------------
# Construccion de parrafos con la estructura XML del deck
# -------------------------------------------------------------------------
def make_para(runs, marL=None, lvl=None, indent=None, bullet=False):
    """Crea un elemento a:p. runs: lista de dicts {text, bold, italic, sz, color}."""
    p = etree.Element(qn("a:p"))
    pPr = etree.SubElement(p, qn("a:pPr"))
    if lvl is not None:
        pPr.set("lvl", str(lvl))
    if marL is not None:
        pPr.set("marL", str(marL))
    if indent is not None:
        pPr.set("indent", str(indent))
    if bullet:
        buFont = etree.SubElement(pPr, qn("a:buFont"))
        buFont.set("typeface", "Courier New")
        buChar = etree.SubElement(pPr, qn("a:buChar"))
        buChar.set("char", "o")
    for spec in runs:
        r = etree.SubElement(p, qn("a:r"))
        rPr = etree.SubElement(r, qn("a:rPr"))
        rPr.set("lang", "es-ES")
        rPr.set("dirty", "0")
        rPr.set("sz", str(spec.get("sz", 1800)))
        if spec.get("bold"):
            rPr.set("b", "1")
        if spec.get("italic"):
            rPr.set("i", "1")
        color = spec.get("color")
        if color is not None:
            sf = etree.SubElement(rPr, qn("a:solidFill"))
            clr = etree.SubElement(sf, qn("a:srgbClr"))
            clr.set("val", color)
        t = etree.SubElement(r, qn("a:t"))
        t.text = spec["text"]
    return p


def clear_paragraphs(txBody_el):
    for p in txBody_el.findall(qn("a:p")):
        txBody_el.remove(p)


def replace_body_with_bullets(txBody_el, bullets):
    """Reemplaza el cuerpo por parrafos con viñeta 'o' (patron del deck)."""
    clear_paragraphs(txBody_el)
    for runs in bullets:
        txBody_el.append(make_para(runs, marL=742950, lvl=1, indent=-285750, bullet=True))


def set_shape_geom(shape, left_in, top_in, width_in, height_in):
    shape.left = Emu(int(left_in * EMU_IN))
    shape.top = Emu(int(top_in * EMU_IN))
    shape.width = Emu(int(width_in * EMU_IN))
    shape.height = Emu(int(height_in * EMU_IN))


def set_shape_geom_el(el, left_in, top_in, width_in, height_in):
    xfrm = el.find(qn("p:spPr")).find(qn("a:xfrm"))
    off = xfrm.find(qn("a:off"))
    ext = xfrm.find(qn("a:ext"))
    off.set("x", str(int(left_in * EMU_IN)))
    off.set("y", str(int(top_in * EMU_IN)))
    ext.set("cx", str(int(width_in * EMU_IN)))
    ext.set("cy", str(int(height_in * EMU_IN)))


# -------------------------------------------------------------------------
# Estimacion de altura de texto (auto-chequeo de desbordes)
# -------------------------------------------------------------------------
def estimate_paragraph_lines(p_el, usable_w_in):
    """Estima (n_lineas_fisicas, altura_promedio) de un a:p.

    Divide el parrafo en lineas logicas (saltos \n dentro de a:t o br) y estima
    para cada una el numero de lineas fisicas segun su tamano de fuente; la altura
    total es la suma de alturas por linea logica (se devuelve altura promedio).
    """
    pPr = p_el.find(qn("a:pPr"))
    mar = 0.0
    if pPr is not None and pPr.get("marL"):
        mar = int(pPr.get("marL")) / EMU_IN
    width = max(1.0, usable_w_in - mar)

    line_spacing = 1.2
    if pPr is not None:
        ln = pPr.find(qn("a:lnSpc"))
        if ln is not None:
            pct = ln.find(qn("a:spcPct"))
            if pct is not None and pct.get("val"):
                line_spacing = int(pct.get("val")) / 100000.0

    # Lineas logicas del parrafo: cada (texto, sz) con su tamano.
    lineas = [[]]
    for r in p_el.iter(qn("a:r")):
        rPr = r.find(qn("a:rPr"))
        sz = 1800  # heredado del tema cuando el run no declara sz
        if rPr is not None and rPr.get("sz"):
            sz = int(rPr.get("sz"))
        t_el = r.find(qn("a:t"))
        texto = t_el.text or "" if t_el is not None else ""
        partes = texto.split("\n")
        for i, parte in enumerate(partes):
            if i > 0:
                lineas.append([])
            if parte:
                lineas[-1].append((parte, sz))
    # Lineas logicas vacias intermedias (doble salto) cuentan como linea vacia;
    # la vacia final (salto al cierre) no genera linea visible.
    while len(lineas) > 1 and not lineas[-1]:
        lineas.pop()

    if not any(lineas):
        return 1, 1800 * 1.2 / 7200.0  # parrafo vacio

    total_h = 0.0
    for linea in lineas:
        if not linea:
            total_h += 1800 * 1.2 / 7200.0
            continue
        sz_max = max(sz for _, sz in linea)
        largo = sum(len(t) for t, _ in linea)
        char_w = 0.5 * (sz_max / 100.0) / 72.0  # ancho medio de caracter Arial
        cpl = max(1, int(width / char_w))
        n_lines = max(1, math.ceil(largo / cpl))
        total_h += n_lines * (sz_max / 100.0) * line_spacing / 72.0
    # Para el auto-chequeo interesa la altura total del parrafo; se devuelve
    # como (n_lineas, altura_promedio) para conservar la firma.
    n_total = max(1, round(total_h / (1800 * line_spacing / 7200.0)))
    return n_total, total_h / n_total if n_total else 0.0


def estimate_text_height(txBody_el, usable_w_in):
    total = 0.0
    for p in txBody_el.findall(qn("a:p")):
        n_lines, line_h = estimate_paragraph_lines(p, usable_w_in)
        total += n_lines * line_h
    return total


# =========================================================================
# Main
# =========================================================================
def main():
    prs = Presentation(CONFIG["deck_companero"])
    assert len(prs.slides) == 17, f"Se esperaban 17 slides, hay {len(prs.slides)}"

    shape_id_start = max_shape_id(prs) + 1
    # Clones: A/D2 <- "Objetivo y pregunta" (idx4); B <- "Corpus" (idx6);
    # C <- "Evolucion del corpus en el tiempo" (idx12); D1 <- section "Conclusiones" (idx14).
    clon_A = clone_slide(prs, 4, shape_id_start)
    clon_B = clone_slide(prs, 6, shape_id_start + 40)
    clon_C = clone_slide(prs, 12, shape_id_start + 80)
    clon_D1 = clone_slide(prs, 14, shape_id_start + 120)
    clon_D2 = clone_slide(prs, 4, shape_id_start + 160)

    # ------------------------------------------------------------
    # Slide A: Contextualizacion / Por que es relevante?
    # ------------------------------------------------------------
    replace_band_title(clon_A, "Contextualización")
    replace_subtitle(clon_A, "¿Por qué es relevante?")
    cuerpo_A = find_shapes(clon_A, name="CuadroTexto 10")[0]
    cuerpo_A.name = "Fusion-A-Cuerpo"
    bullets_A = [
        "El análisis manual de decenas de miles de documentos largos es inviable: el PLN "
        "convierte ese contenido en evidencia cuantificable y reproducible.",
        "Escasez de recursos y datasets legales anotados en español para América Latina: un "
        "corpus jurídico colombiano depurado y caracterizado textométricamente constituye "
        "un aporte en sí mismo.",
        "Línea base para clasificación temática de decisiones, seguimiento de la agenda de la "
        "Corte y consulta de jurisprudencia; siempre como apoyo a la interpretación "
        "especializada, no como sustituto.",
    ]
    replace_body_with_bullets(cuerpo_A._element.find(qn("p:txBody")),
                              [[{"text": b}] for b in bullets_A])
    est = estimate_text_height(cuerpo_A._element.find(qn("p:txBody")),
                               cuerpo_A.width / EMU_IN)
    set_shape_geom(cuerpo_A, 0.75, 2.694, 11.42,
                   min(max(2.6, est + 0.35), 7.38 - 2.694))

    # ------------------------------------------------------------
    # Slide B: Construccion del corpus / Fragmentos representativos
    # ------------------------------------------------------------
    replace_band_title(clon_B, "Construcción del corpus")
    replace_subtitle(clon_B, "Fragmentos representativos")
    for shp in find_shapes(clon_B, name="CuadroTexto 4"):
        remove_shape(clon_B, shp)
    # Tarjetas: mismo estilo que la tarjeta INGESTA de la slide 8 (roundRect,
    # fill EAF2F8, borde 2F75B5 1.5pt, radio de esquina por defecto de la plantilla).
    src_card = prs.slides[7].shapes[3]._element
    citas = [
        ("Sentencia T-293/97",
         "La acción de tutela es mecanismo subsidiario cuyo objeto específico es la "
         "protección de los derechos fundamentales violados o amenazados por acción u "
         "omisión de una autoridad pública…"),
        ("Sentencia C-448/98",
         "La Corte carece de competencia para conocer sobre la constitucionalidad de las "
         "preceptivas que integran el texto normativo contenido en el Decreto 1421 de 1993…"),
        ("Sentencia SU-476/97",
         "La vida en comunidad conlleva forzosamente el cumplimiento de una serie de deberes "
         "recíprocos por parte de los asociados, el primero de los cuales es el de respetar "
         "los derechos de los demás."),
    ]
    left_card, w_card = 0.75, 11.83
    top_card, h_card, gap_card = 2.56, 1.30, 0.21
    for k, (iden, cita) in enumerate(citas):
        card = copy.deepcopy(src_card)
        set_shape_geom_el(card, left_card, top_card + k * (h_card + gap_card),
                          w_card, h_card)
        cnv = card.find(qn("p:nvSpPr")).find(qn("p:cNvPr"))
        cnv.set("name", f"Fusion-B-Tarjeta{k + 1}")
        txBody = card.find(qn("p:txBody"))
        clear_paragraphs(txBody)
        bodyPr = txBody.find(qn("a:bodyPr"))
        if bodyPr is not None:
            bodyPr.set("anchor", "t")
        txBody.append(make_para(
            [{"text": iden, "bold": True, "sz": 1400, "color": AZUL}], marL=0, lvl=0))
        txBody.append(make_para(
            [{"text": cita, "italic": True, "sz": 1100, "color": GRIS_TARJETA}],
            marL=0, lvl=0))
        clon_B.shapes._spTree.append(card)
    # Nota inferior
    tb_nota_B = clon_B.shapes.add_textbox(
        Emu(int(1.0 * EMU_IN)),
        Emu(int((top_card + 3 * h_card + 2 * gap_card + 0.05) * EMU_IN)),
        Emu(int(11.33 * EMU_IN)), Emu(int(0.32 * EMU_IN)))
    tb_nota_B.name = "Fusion-B-Nota"
    tf = tb_nota_B.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = 2  # CENTER
    r = p.add_run()
    r.text = "Fragmentos literales del anexo del corpus (muestra de 25 sentencias)."
    r.font.name = "Arial"
    r.font.size = Pt(11)
    r.font.italic = True
    r.font.color.rgb = _rgb(GRIS_TEXTO)

    # ------------------------------------------------------------
    # Slide C: Textometria / Nubes de palabras por anio
    # ------------------------------------------------------------
    replace_band_title(clon_C, "Textometría")
    replace_subtitle(clon_C, "Nubes de palabras por año")
    # Quitar solo la grafica ancha del patron (la banda 13.33x1.77 y el icono se conservan).
    for shp in clon_C.shapes:
        if (shp.shape_type == 13 and shp.width / EMU_IN > 4.0
                and shp.height / EMU_IN > 3.0):
            remove_shape(clon_C, shp)
    dims = []
    for rango, path in CONFIG["nubes"].items():
        with Image.open(path) as im:
            ratio = im.width / im.height
        dims.append((rango, path, CONFIG["altura_nube_in"] * ratio))
    total_w = sum(w for _, _, w in dims) + CONFIG["gap_nubes_in"] * (len(dims) - 1)
    x0 = (13.333 - total_w) / 2.0
    y_img = 2.47
    etiquetas = {"1992-2001": "1992–2001", "2002-2011": "2002–2011", "2012-2020": "2012–2020"}
    cur_x = x0
    for rango, path, w in dims:
        pic = clon_C.shapes.add_picture(
            path, Emu(int(cur_x * EMU_IN)), Emu(int(y_img * EMU_IN)),
            Emu(int(w * EMU_IN)), Emu(int(CONFIG["altura_nube_in"] * EMU_IN)))
        pic.name = f"Fusion-C-Nube-{rango}"
        # Etiqueta bajo la imagen (misma anchura, centrada)
        tb = clon_C.shapes.add_textbox(
            Emu(int(cur_x * EMU_IN)), Emu(int((y_img + CONFIG["altura_nube_in"] + 0.05) * EMU_IN)),
            Emu(int(w * EMU_IN)), Emu(int(0.3 * EMU_IN)))
        tb.name = f"Fusion-C-Etiqueta-{rango}"
        tf = tb.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.alignment = 2
        r = p.add_run()
        r.text = etiquetas[rango]
        r.font.name = "Arial"
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = _rgb(AZUL)
        cur_x += w + CONFIG["gap_nubes_in"]
    # Nota tecnica
    tb_notaC = clon_C.shapes.add_textbox(
        Emu(int(0.75 * EMU_IN)), Emu(int(6.97 * EMU_IN)),
        Emu(int(11.83 * EMU_IN)), Emu(int(0.35 * EMU_IN)))
    tb_notaC.name = "Fusion-C-Nota"
    tf = tb_notaC.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = 2
    r = p.add_run()
    r.text = ("Top-1.000 términos por año; sin stopwords gramaticales (306 formas) ni de "
              "dominio jurídico (65 formas). Año 2021 excluido (<50 documentos).")
    r.font.name = "Arial"
    r.font.size = Pt(11)
    r.font.italic = True
    r.font.color.rgb = _rgb(GRIS_TEXTO)

    # ------------------------------------------------------------
    # Slide D1: header de seccion "Antecedentes"
    # ------------------------------------------------------------
    for shp in clon_D1.shapes:
        if not shp.has_text_frame:
            continue
        for p in shp.text_frame.paragraphs:
            for r in p.runs:
                if r.text.strip() == "Conclusiones":
                    set_run_text(r, "Antecedentes")
                    break

    # ------------------------------------------------------------
    # Slide D2: Antecedentes / Estado del arte
    # ------------------------------------------------------------
    replace_band_title(clon_D2, "Antecedentes")
    replace_subtitle(clon_D2, "Estado del arte")
    cuerpo_D2 = find_shapes(clon_D2, name="CuadroTexto 10")[0]
    cuerpo_D2.name = "Fusion-D2-Cuerpo"
    refs = [
        ("Anzola Cortés (2021): ",
         "minería de datos sobre más de 6.000 sentencias de constitucionalidad de la Corte "
         "Constitucional colombiana (1992-2020): el mismo universo documental del corpus."),
        ("Barbosa Becerra (2021): ",
         "análisis léxico-discursivo de sentencias de la misma Corte; evidencia fórmulas "
         "discursivas estables y léxico especializado."),
        ("HaCohen-Kerner et al. (2020): ",
         "la elección del preprocesamiento (minúsculas, números, stopwords, normalización) "
         "cambia los resultados según el corpus y la tarea: por eso esta entrega documenta "
         "cada paso."),
        ("González-González et al. (2023): ",
         "clasificación automática de sentencias judiciales en español con diccionarios de "
         "dominio construidos con expertos legales."),
        ("Chalkidis et al. (2020) (LEGAL-BERT) y Raj et al. (2026): ",
         "vocabulario jurídico distinto del general y cambios léxicos medibles en el tiempo "
         "motivan la adaptación de dominio y el análisis diacrónico en entregas siguientes."),
    ]
    bullets_D2 = [[{"text": autor, "bold": True}, {"text": detalle}]
                  for autor, detalle in refs]
    txBody_D2 = cuerpo_D2._element.find(qn("p:txBody"))
    replace_body_with_bullets(txBody_D2, bullets_D2)
    txBody_D2.append(make_para(
        [{"text": "No hay trabajos indexados que usen el dataset exacto de Hugging Face "
                  "(verificado 2026-09-05); los antecedentes cubren el corpus subyacente y "
                  "la metodología.", "italic": True, "sz": 1200, "color": GRIS_TEXTO}],
        marL=457200, lvl=1))
    est = estimate_text_height(txBody_D2, cuerpo_D2.width / EMU_IN)
    set_shape_geom(cuerpo_D2, 0.75, 2.694, 11.42,
                   min(max(2.6, est + 0.35), 7.38 - 2.694))

    # ------------------------------------------------------------
    # Edicion E: paso 09 (slide 10 original, indice 9)
    # ------------------------------------------------------------
    paso09 = None
    for shp in prs.slides[9].shapes:
        if shp.has_text_frame and "09  Eliminar stopwords" in shp.text_frame.text:
            paso09 = shp
            break
    assert paso09 is not None, "No se encontro el paso 09"
    paso09.name = "Fusion-E-Paso09"
    for p in paso09.text_frame.paragraphs:
        for r in p.runs:
            if "Conservar términos informativos" in r.text:
                set_run_text(r, "\nConservar términos informativos · "
                                "306 formas normalizadas (sin tildes)")
                # La caja mide 1.0in y el detalle gana una linea: se reduce 1pt
                # (12 -> 11) manteniendo legibilidad (misma familia que el resto).
                rpr = r._r.find(qn("a:rPr"))
                if rpr is not None:
                    rpr.set("sz", "1100")

    # ------------------------------------------------------------
    # Edicion F: dos lineas al final de la ficha tecnica (slide 7 original, indice 6)
    # ------------------------------------------------------------
    ficha = find_shapes(prs.slides[6], name="CuadroTexto 4")[0]
    ficha.name = "Fusion-F-Ficha"
    txBody_F = ficha._element.find(qn("p:txBody"))
    txBody_F.append(make_para(
        [{"text": "Formato de almacenamiento: ", "bold": True},
         {"text": "CSV UTF-8 (1,68 GB), dos columnas (identificador + texto completo)."}],
        marL=457200, lvl=1))
    txBody_F.append(make_para(
        [{"text": "Nota: ", "bold": True},
         {"text": "el archivo del repositorio se llama Juris1992_2022.csv, pero la cobertura "
                  "real es 1992-2021 (2021 parcial: 20 sentencias)."}],
        marL=457200, lvl=1))
    paras = txBody_F.findall(qn("a:p"))
    if paras and not "".join(t.text or "" for t in paras[-1].iter(qn("a:t"))).strip():
        txBody_F.remove(paras[-1])
    # El texto crece ~4 lineas; para que quepa en el area disponible (0.75-7.2in)
    # se compacta ligeramente el interlineado de la ficha (92 %) manteniendo
    # legibilidad, segun lo permitido por la consigna.
    for p_el in txBody_F.findall(qn("a:p")):
        pPr = p_el.find(qn("a:pPr"))
        if pPr is None:
            pPr = etree.SubElement(p_el, qn("a:pPr"))
            pPr.text = None
            p_el.insert(0, pPr)
        ln = etree.SubElement(pPr, qn("a:lnSpc"))
        sp = etree.SubElement(ln, qn("a:spcPct"))
        sp.set("val", "92000")
    est = estimate_text_height(txBody_F, ficha.width / EMU_IN)
    top_ficha = ficha.top / EMU_IN
    ficha.height = Emu(int(max(ficha.height / EMU_IN, min(est + 0.3, 7.2 - top_ficha)) * EMU_IN))

    # ------------------------------------------------------------
    # Edicion G: agenda (slide 2 original, indice 1)
    # ------------------------------------------------------------
    update_agenda(prs, prs.slides[1])

    # ------------------------------------------------------------
    # Reordenar slides al orden final
    # ------------------------------------------------------------
    reorder_slides(prs, clon_A, clon_B, clon_C, clon_D1, clon_D2)

    prs.save(CONFIG["deck_companero"])
    print("OK: guardado en", CONFIG["deck_companero"])


# =========================================================================
# Agenda (edicion G)
# =========================================================================
def update_agenda(prs, slide_agenda):
    """Renumeracion de las 5 secciones con los rangos reales finales:
    01 Contextualizacion 03–06 | 02 Construccion del corpus 07–10 |
    03 Textometria 11–17 | 04 Antecedentes 18–19 | 05 Conclusiones 20–21."""
    filas = agenda_rows(slide_agenda)
    assert len(filas) == 4, f"Agenda con {len(filas)} filas, se esperaban 4"

    # Posiciones verticales finales de las 5 filas (pitch 1.0in, alto 0.889in).
    y_final = [2.472, 3.472, 4.472, 5.472, 6.472]
    y_orig = [2.472, 3.556, 4.639, 5.722]

    # Fila 5: clon de la fila 4 original (Conclusiones, colores purpura).
    fila5 = [copy.deepcopy(el) for el in filas[3]]
    delta5 = y_final[4] - y_orig[3]
    for el in fila5:
        offset_y(el, delta5)
        rename_cnv(el, "Fusion-G-Fila5")
    for el in fila5:
        slide_agenda.shapes._spTree.append(el)

    # Reposicionar filas 2, 3 y 4 (la fila 1 no se mueve).
    for idx in (1, 2, 3):
        delta = y_final[idx] - y_orig[idx]
        for el in filas[idx]:
            offset_y(el, delta)

    # Contenido de cada fila (el orden XML dentro de fila es: rect, num, nombre, rango).
    specs = [
        ("01", "Contextualización", "Reto, justificación, objetivo y preguntas", "03–06", None),
        ("02", "Construcción del corpus", "Fuente, depuración, fragmentos y validación",
         "07–10", None),
        ("03", "Textometría", "Procesamiento, dimensiones y resultados léxicos", "11–17", None),
        ("04", "Antecedentes", "Estado del arte", "18–19", (AZUL, AZUL_CLARO)),
        ("05", "Conclusiones", "Hallazgos y siguientes etapas", "20–21", None),
    ]
    fila_objs = [filas[0], filas[1], filas[2], filas[3], fila5]
    for fila, (num, nombre, detalle, rango, recolor) in zip(fila_objs, specs):
        if recolor is not None:
            recolor_fila(fila, recolor[0], recolor[1])
        set_fila_texto(fila, num, nombre, detalle, rango)


def agenda_rows(slide_agenda):
    """4 filas de la agenda; cada fila = [rect(roundRect sin texto), num, nombre, rango]."""
    filas = []
    actual = []
    for shp in slide_agenda.shapes:
        is_row_rect = (shp.shape_type == 1
                       and shp.name.startswith("Rectángulo: esquinas redondeadas")
                       and not (shp.has_text_frame and shp.text_frame.text.strip()))
        if is_row_rect:
            if actual:
                filas.append(actual)
            actual = [shp._element]
        elif actual:
            actual.append(shp._element)
    if actual:
        filas.append(actual)
    return filas


def offset_y(el, delta_in):
    off = el.find(qn("p:spPr")).find(qn("a:xfrm")).find(qn("a:off"))
    off.set("y", str(int(int(off.get("y")) + delta_in * EMU_IN)))


def rename_cnv(el, prefix):
    cnv = el.find(qn("p:nvSpPr")).find(qn("p:cNvPr"))
    cnv.set("name", prefix + "-" + cnv.get("name", "")[:25])


def set_fila_texto(fila, num, nombre, detalle, rango):
    """Escribe el contenido de una fila: elementos [rect, num, nombre+detalle, rango]."""
    el_num, el_nombre, el_rango = fila[1], fila[2], fila[3]
    for t in el_num.iter(qn("a:t")):
        t.text = num
    runs = list(el_nombre.iter(qn("a:r")))
    if runs:
        t0 = runs[0].find(qn("a:t"))
        if t0 is not None:
            t0.text = nombre
        if len(runs) >= 2:
            t1 = runs[1].find(qn("a:t"))
            if t1 is not None:
                t1.text = "\n" + detalle
        else:
            r1 = etree.SubElement(runs[0].getparent(), qn("a:r"))
            rPr = etree.SubElement(r1, qn("a:rPr"))
            rPr.set("lang", "es-ES")
            rPr.set("sz", "1200")
            sf = etree.SubElement(rPr, qn("a:solidFill"))
            clr = etree.SubElement(sf, qn("a:srgbClr"))
            clr.set("val", GRIS_TEXTO)
            tt = etree.SubElement(r1, qn("a:t"))
            tt.text = "\n" + detalle
            runs[0].getparent().insert(1, r1)
    for t in el_rango.iter(qn("a:t")):
        t.text = rango


def recolor_fila(fila, accent, fill):
    rect = fila[0]
    for c in rect.iter(qn("a:srgbClr")):
        if c.get("val") == LILA:
            c.set("val", fill)
        elif c.get("val") == PURPURA:
            c.set("val", accent)
    for el in fila[1:]:
        for c in el.iter(qn("a:srgbClr")):
            if c.get("val") == PURPURA:
                c.set("val", accent)


# =========================================================================
# Orden final de slides
# =========================================================================
def reorder_slides(prs, clon_A, clon_B, clon_C, clon_D1, clon_D2):
    clones = (clon_A, clon_B, clon_C, clon_D1, clon_D2)
    origs = [s for s in prs.slides if not any(s is c for c in clones)]
    assert len(origs) == 17, f"Hay {len(origs)} slides originales"
    orden = []
    for idx, s in enumerate(origs):
        orden.append(s)
        if idx == 4:    # tras "Objetivo y pregunta de investigación" -> A
            orden.append(clon_A)
        if idx == 7:    # tras "Depuración y enriquecimiento del corpus" -> B
            orden.append(clon_B)
        if idx == 13:   # tras "Más allá del vocabulario jurídico": C + seccion Antecedentes
            orden.append(clon_C)
            orden.append(clon_D1)
            orden.append(clon_D2)
    assert len(orden) == 22, f"Orden final con {len(orden)} slides"
    sldIdLst = prs.slides._sldIdLst
    pres_part = prs.part
    mapa = {}
    for sldId in sldIdLst:
        rid = sldId.get(qn("r:id"))
        part = pres_part.related_part(rid)
        mapa[str(part.partname)] = sldId
    for sldId in list(sldIdLst):
        sldIdLst.remove(sldId)
    for s in orden:
        sldIdLst.append(mapa[str(s.part.partname)])


# =========================================================================
# Verificacion
# =========================================================================
def verificar():
    prs = Presentation(CONFIG["deck_companero"])
    n = len(prs.slides)
    assert n == 22, f"Se esperaban 22 slides, hay {n}"
    print(f"\nVERIFICACION: {n} slides totales")
    for i, s in enumerate(prs.slides, 1):
        txts, n_pics = [], 0
        for sh in s.shapes:
            if sh.shape_type == 13:
                n_pics += 1
            if sh.has_text_frame and sh.text_frame.text.strip():
                txts.append(sh.text_frame.text.strip().split("\n")[0][:60])
        print(f"  slide {i:2d} | pics={n_pics:2d} | {txts[0] if txts else '(sin texto)'}")
    slide17 = prs.slides[16]
    nubes = [sh for sh in slide17.shapes
             if sh.shape_type == 13 and 2.0 < sh.width / EMU_IN < 6.0
             and sh.height / EMU_IN > 3.0]
    assert len(nubes) == 3, f"La slide 17 debe tener las 3 nubes anuales, tiene {len(nubes)}"
    print(f"  -> slide 17 OK: 3 nubes anuales (anchos "
          f"{[round(sh.width / EMU_IN, 2) for sh in nubes]}in)")


def auto_chequeo_desbordes():
    """Cajas nuevas/editadas: geometria dentro de 13.333x7.4 y texto estimado <= caja."""
    prs = Presentation(CONFIG["deck_companero"])
    problemas = []
    revisadas = 0
    for i, s in enumerate(prs.slides, 1):
        for sh in s.shapes:
            if not sh.name.startswith("Fusion-"):
                continue  # solo cajas nuevas/editadas por la fusion
            if not (sh.has_text_frame and sh.text_frame.text.strip()):
                continue
            left, top = sh.left / EMU_IN, sh.top / EMU_IN
            width, height = sh.width / EMU_IN, sh.height / EMU_IN
            # 1) Limite geometrico: dentro de 13.333 x 7.4 in.
            if (left < -0.02 or top < -0.02 or left + width > 13.333 + 0.02
                    or top + height > 7.4 + 0.02):
                problemas.append(
                    f"slide {i} '{sh.name}': geometria L={left:.2f} T={top:.2f} "
                    f"R={left + width:.2f} B={top + height:.2f}")
                continue
            # 2) Desborde estimado del texto respecto de la caja.
            est = estimate_text_height(sh._element.find(qn("p:txBody")), width)
            if est > height + 0.12:
                problemas.append(
                    f"slide {i} '{sh.name}': texto estimado {est:.2f}in > caja "
                    f"{height:.2f}in (B={top + height:.2f})")
            else:
                revisadas += 1
    print(f"\nAUTO-CHEQUEO DESBORDES: {revisadas} cajas nuevas/editadas OK")
    if problemas:
        print("  PROBLEMAS DETECTADOS:")
        for p in problemas:
            print("   -", p)
    else:
        print("  Resultado: TODO OK (ninguna caja fuera de 13.333x7.4 ni desbordada)")
    return not problemas


def verificar_agenda_rangos():
    """Cada rango NN–MM de la agenda corresponde a slides reales con la seccion esperada."""
    prs = Presentation(CONFIG["deck_companero"])
    esperados = [
        ("03–06", "Contextualización", 6),
        ("07–10", "Construcción del corpus", 10),
        ("11–17", "Textometría", 17),
        ("18–19", "Antecedentes", 19),
        ("20–21", "Conclusiones", 21),
    ]
    print("\nVERIFICACION DE RANGOS DE LA AGENDA (slide inicial = header de seccion "
          "o primera slide de la seccion):")
    ok = True
    for rango, seccion, ultima in esperados:
        ini, fin = (int(x) for x in rango.split("–"))
        assert fin == ultima, f"Rango {rango} no coincide con el orden real"
        s_ini = prs.slides[ini - 1]
        tit = next((sh.text_frame.text.strip().split("\n")[0]
                    for sh in s_ini.shapes
                    if sh.has_text_frame and sh.text_frame.text.strip()), "")
        if seccion not in tit:
            ok = False
            print(f"  {rango} {seccion}: PROBLEMA (slide {ini} dice '{tit}')")
        else:
            print(f"  {rango} {seccion}: OK (slide {ini} -> '{tit}', ultima {fin})")
    return ok


if __name__ == "__main__":
    try:
        main()
        verificar()
        auto_chequeo_desbordes()
        verificar_agenda_rangos()
        print("\nFin de la ejecucion.")
    except Exception:
        traceback.print_exc()
        sys.exit(1)
