#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Validacion de consistencia de datos entre el notebook (metricas.json),
las diapositivas (PPTX) y el informe (main.tex / PDF).

Fuente de verdad: resultados/metricas.json (generado por el notebook) mas el
total sin stopwords reportado por el propio notebook (128.806.664).
Uso: /home/carlos/venv/bin/python scripts/validar_datos.py
Salida: reporte plan/validacion_datos.md y resumen en consola.
"""
import json
import os
import re
import subprocess
import sys
import zipfile

from pptx import Presentation

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

ARTEFACTOS_TEXTO = {
    "deck_companero": os.path.join(BASE, "PLN_PROYECTO.pptx"),
    "deck_nuestro": os.path.join(BASE, "diapositivas", "Entrega1_Diapositivas.pptx"),
    "informe_tex": os.path.join(BASE, "informe", "main.tex"),
    "informe_pdf": os.path.join(BASE, "Entrega_1_Informe.pdf"),
}

TOL_REL = 0.0015  # 0,15 %: redondeos presentacionales de diapositivas
EXACTOS_INFORME = [  # cifras que en el informe deben aparecer EXACTAS (tabla/texto)
    "23.742", "17.099", "6.246", "397", "260.988.316", "250.012.101",
    "128.806.664", "10.992,68", "10.530,37", "947.616", "898.926", "809.313",
    "760.197", "736.326", "22.543.375", "99,97", "306", "48",
]


def texto_pptx(ruta):
    """Extrae todo el texto de las diapositivas."""
    prs = Presentation(ruta)
    partes = []
    for slide in prs.slides:
        for sh in slide.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                partes.append(sh.text_frame.text)
            if sh.shape_type == 6:  # grupo
                try:
                    for sub in sh.shapes:
                        if sub.has_text_frame and sub.text_frame.text.strip():
                            partes.append(sub.text_frame.text)
                except Exception:
                    pass
    return "\n".join(partes)


def texto_pdf(ruta):
    out = subprocess.run(["pdftotext", "-layout", ruta, "-"],
                         capture_output=True, text=True)
    return out.stdout if out.returncode == 0 else ""


def numeros_normalizados(texto):
    """Devuelve todos los numeros (flotantes) presentes en un texto."""
    limpio = re.sub(r"[.,](?=\d{3}(?:\D|$))", "", texto)  # miles es
    limpio = limpio.replace(",", ".").replace("≈", "").replace("M", "")
    return [float(x) for x in re.findall(r"\d+(?:\.\d+)?", limpio)]


def carga_canonica():
    with open(os.path.join(BASE, "resultados", "metricas.json")) as f:
        m = json.load(f)
    tipos = m["distribucion_tipos"]
    canon = {
        "n_docs_total_csv": m["n_docs_total_csv"],
        "n_docs_parseados": m["n_docs_parseados"],
        "n_docs_excluidos": m["n_docs_excluidos"],
        "T": tipos["T"], "C": tipos["C"], "SU": tipos["SU"],
        "total_raw": m["total_palabras_raw"],
        "total_clean": m["total_palabras_clean"],
        "total_sin_stopwords": 128_806_664,
        "prom_raw": m["promedio_palabras_raw_por_doc"],
        "prom_clean": m["promedio_palabras_clean_por_doc"],
        "annio_min": 1992, "annio_max": 2021,
        "docs_2021": m["docs_por_anio"].get("2021", 20),
        "pct_stopwords": 48.0,
        "tasa_parseo": 99.97,
        "excluidos_autos": 6, "excluidos_html": 2,
        "anexo": 25, "top1_con": 22_543_375, "top1_sin": 947_616,
        "stopwords_formas": 306,
    }
    return canon


def main():
    canon = carga_canonica()
    textos = {}
    for nombre, ruta in ARTEFACTOS_TEXTO.items():
        if not os.path.exists(ruta):
            textos[nombre] = ""
            continue
        if ruta.endswith(".pptx"):
            textos[nombre] = texto_pptx(ruta)
        elif ruta.endswith(".pdf"):
            textos[nombre] = texto_pdf(ruta)
        else:
            with open(ruta, encoding="utf-8") as f:
                textos[nombre] = f.read()

    lineas = ["# Validacion de consistencia de datos (Fase 7, Tarea 0)",
              "",
              "Fuente de verdad: `resultados/metricas.json` (notebook, corpus completo) y total sin stopwords reportado por el notebook (128.806.664).",
              "Tolerancia en diapositivas: redondeos presentacionales con error relativo < 0,15 %.",
              "En el informe (tex/pdf) las cifras de la tabla deben ser exactas.",
              "",
              "| Cifra (canonica) | Deck compañero | Deck nuestro | Informe tex | Informe pdf |",
              "|---|---|---|---|---|"]

    problemas = []
    for clave, valor in canon.items():
        filas = []
        for nombre in ("deck_companero", "deck_nuestro", "informe_tex", "informe_pdf"):
            texto = textos[nombre]
            numeros = numeros_normalizados(texto)
            ok = any(abs(n - valor) / valor <= TOL_REL for n in numeros)
            # busqueda exacta de la cifra con formato espanol
            s_es = f"{valor:,}".replace(",", ".")
            ok = ok or (s_es in texto)
            # millones con coma (p.ej. 260,99 M) -> formato es: 260,99
            if valor >= 1e6:
                s_mil = f"{valor / 1e6:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
                ok = ok or any(v in texto for v in (s_mil, s_mil.replace(",", ".")))
            filas.append("OK" if ok else "FALTA")
        estado = "ok" if filas.count("FALTA") == 0 else "REVISAR"
        if estado == "REVISAR":
            problemas.append(clave)
        # nombre legible
        nombre_cifra = {
            "n_docs_total_csv": "Filas CSV (23.750)", "n_docs_parseados": "Sentencias (23.742)",
            "n_docs_excluidos": "Excluidas (8)", "T": "Tutela (17.099)", "C": "Constitucionalidad (6.246)",
            "SU": "Unificacion (397)", "total_raw": "Palabras raw (260.988.316)",
            "total_clean": "Palabras limpias (250.012.101)", "total_sin_stopwords": "Sin stopwords (128.806.664)",
            "prom_raw": "Promedio raw/doc (10.992,68)", "prom_clean": "Promedio limpio/doc (10.530,37)",
            "annio_min": "Anio minimo (1992)", "annio_max": "Anio maximo (2021)",
            "docs_2021": "Docs 2021 (20)", "pct_stopwords": "Stopwords (48 %)", "tasa_parseo": "Parseo (99,97 %)",
            "excluidos_autos": "Excluidos autos (6)", "excluidos_html": "Excluidos HTML (2)",
            "anexo": "Anexo (25)", "top1_con": "Top-1 con (de 22.543.375)",
            "top1_sin": "Top-1 sin (articulo 947.616)", "stopwords_formas": "Stopwords formas (306)",
        }.get(clave, clave)
        lineas.append("| {} | {} | {} | {} | {} |".format(nombre_cifra, *filas))

    # Chequeo de cifras exactas en el informe (tex)
    texto_tex = textos["informe_tex"]
    exactas_faltantes = [s for s in EXACTOS_INFORME if s not in texto_tex]
    if exactas_faltantes:
        problemas.append("cifras_exactas_informe")
    lineas.append("")
    lineas.append("Cifras exactas exigidas al informe (tex) que NO aparecen: {}"
                  .format(", ".join(exactas_faltantes) if exactas_faltantes else "ninguna (todas presentes)"))
    lineas.append("")

    # md5 de las figuras del deck del companero vs figuras del notebook
    lineas.append("## Figuras incrustadas vs figuras del notebook (md5)")
    import hashlib

    def md5_archivo(ruta):
        h = hashlib.md5()
        with open(ruta, "rb") as f:
            for bloque in iter(lambda: f.read(1 << 20), b""):
                h.update(bloque)
        return h.hexdigest()

    figs_notebook = {
        "fig_top20_con_stopwords.png": os.path.join(BASE, "figuras", "fig_top20_con_stopwords.png"),
        "fig_top20_sin_stopwords.png": os.path.join(BASE, "figuras", "fig_top20_sin_stopwords.png"),
        "fig_sentencias_por_anio.png": os.path.join(BASE, "figuras", "fig_sentencias_por_anio.png"),
        "fig_nube_global.png": os.path.join(BASE, "figuras", "fig_nube_global.png"),
    }
    hashes_nb = {k: md5_archivo(v) for k, v in figs_notebook.items()}
    con_zip = zipfile.ZipFile(ARTEFACTOS_TEXTO["deck_companero"])
    media = {n: hashlib.md5(con_zip.read(n)).hexdigest()
             for n in con_zip.namelist() if n.startswith("ppt/media/")}
    con_zip.close()
    for nombre, h in hashes_nb.items():
        en_deck = any(h == mh for mh in media.values())
        lineas.append("- {}: {} en el deck del companero".format(
            nombre, "presente (md5 identico)" if en_deck else "NO presente"))
    lineas.append("")
    with open(os.path.join(BASE, "plan", "validacion_datos.md"), "w", encoding="utf-8") as f:
        f.write("\n".join(lineas))

    print("Cifras validadas:", len(canon))
    print("Cifras con FALTA en algun artefacto:", problemas if problemas else "ninguna")
    print("Reporte guardado en plan/validacion_datos.md")


if __name__ == "__main__":
    main()
