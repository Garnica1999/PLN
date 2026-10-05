"""generar_presentacion.py

Fecha: 2026-10-04
Descripcion: genera la presentacion PowerPoint (16:9) de la Entrega 2 del
proyecto PLN (EDA avanzado, representacion de textos y propuesta
metodologica) a partir de las cifras verificadas del informe final
(informe/main.tex; resultados/metricas_entrega_ii.json) y de las figuras
PNG existentes en figuras/. La salida es
presentacion/Entrega_2_Presentacion.pptx (12 diapositivas).

Diseno (decisiones documentadas):
- Slide size 16:9 (13.333 x 7.5 in), plantilla en blanco.
- Banda de titulo azul oscuro #1F3864 con texto blanco 26 pt bold.
- Texto de cuerpo #333333; bullets 13-15 pt; paneles con fondo claro #EEF3FA.
- Figuras insertadas con add_picture conservando la proporcion real del PNG
  (dimensiones leidas con PIL); 8 figuras integradas (diapositivas 5, 6, 7 y 10).
- Notas del orador en las 12 diapositivas; numero de diapositiva en el pie
  (a partir de la 2).
- Espera con reintentos si una figura esta siendo regenerada por el notebook
  (mismos valores; el archivo puede no estar completo durante la escritura).

Uso:
    /home/carlos/venv/bin/python presentacion/generar_presentacion.py

El contenido reproduce las cifras verificadas del informe de la Entrega 2;
no se inventan valores.
"""

from __future__ import annotations

import time
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

# ---------------------------------------------------------------------------
# Paleta y parametros de diseno
# ---------------------------------------------------------------------------
DARK_BLUE = RGBColor(0x1F, 0x38, 0x64)
SECONDARY_BLUE = RGBColor(0x2F, 0x5D, 0x8A)
TEXT_COLOR = RGBColor(0x33, 0x33, 0x33)
GRAY_COLOR = RGBColor(0x59, 0x59, 0x59)
LIGHT_GRAY = RGBColor(0x8A, 0x8A, 0x8A)
LIGHT_BG = RGBColor(0xEE, 0xF3, 0xFA)
NOTE_BG = RGBColor(0xF2, 0xF5, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
FONT = "Calibri"

BASE_DIR = Path(__file__).resolve().parents[1]
FIG_DIR = BASE_DIR / "figuras"
OUT_DIR = BASE_DIR / "presentacion"
OUT_FILE = OUT_DIR / "Entrega_2_Presentacion.pptx"

# Figuras usadas en la presentacion (8 en total)
FIG_POR_ANIO = FIG_DIR / "fig_sentencias_por_anio.png"
FIG_TOP20_SIN = FIG_DIR / "fig_top20_sin_stopwords.png"
FIG_NGRAMAS = FIG_DIR / "fig_ngramas_bi_tri.png"
FIG_COLOCACIONES = FIG_DIR / "fig_colocaciones.png"
FIG_MTLD = FIG_DIR / "fig_mtld_distribucion.png"
FIG_LEMMA_STEM = FIG_DIR / "fig_lemma_vs_stem.png"
FIG_TFIDF = FIG_DIR / "fig_representacion_tfidf.png"
FIG_EMBEDDINGS = FIG_DIR / "fig_embeddings_ejemplo.png"

FIGURAS_USADAS = [
    FIG_POR_ANIO, FIG_TOP20_SIN, FIG_NGRAMAS, FIG_COLOCACIONES,
    FIG_MTLD, FIG_LEMMA_STEM, FIG_TFIDF, FIG_EMBEDDINGS,
]


# ---------------------------------------------------------------------------
# Utilidades de figuras (con reintentos si el notebook las regenera)
# ---------------------------------------------------------------------------
def esperar_figura(path: Path, intentos: int = 8, espera_seg: float = 3.0) -> None:
    """Espera a que un PNG exista y sea legible.

    El notebook de la entrega puede estar regenerando las figuras en
    paralelo; este guard reintenta antes de fallar.

    Args:
        path: Ruta del PNG.
        intentos: Numero maximo de intentos.
        espera_seg: Segundos de espera entre intentos.

    Raises:
        FileNotFoundError: Si el archivo no queda disponible.
    """
    for intento in range(1, intentos + 1):
        if path.exists():
            try:
                with Image.open(path) as im:
                    im.verify()
                return
            except Exception:  # noqa: BLE001 - archivo en escritura; se reintenta
                pass
        if intento < intentos:
            time.sleep(espera_seg)
    raise FileNotFoundError(
        "Figura no disponible tras {} intentos: {}".format(intentos, path))


def png_size(path: Path) -> tuple[int, int]:
    """Devuelve (ancho, alto) en pixeles de un PNG usando PIL."""
    with Image.open(path) as im:
        return im.size


# ---------------------------------------------------------------------------
# Utilidades de texto
# ---------------------------------------------------------------------------
def est_lineas(texto: str, size_pt: float, width_in: float) -> int:
    """Estimacion conservadora de lineas que ocupara un texto al envolverse."""
    if not texto:
        return 0
    ancho_char = 0.55 * size_pt / 72.0
    chars_por_linea = max(1, int(width_in / ancho_char))
    total = 0
    for segmento in texto.split("\n"):
        total += max(1, -(-len(segmento) // chars_por_linea))
    return total


def agregar_parrafo(tf, texto, size=14, bold=False, color=TEXT_COLOR,
                    space_after=6, space_before=0, align=PP_ALIGN.LEFT,
                    bullet=None, italic=False, first=False, line_spacing=1.0):
    """Agrega un parrafo con un run unico al text frame dado."""
    p = tf.paragraphs[0] if first and not tf.paragraphs[0].runs else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    p.space_before = Pt(space_before)
    if line_spacing:
        p.line_spacing = line_spacing
    run = p.add_run()
    texto_final = texto
    if bullet:
        texto_final = bullet + texto_final
    run.text = texto_final
    f = run.font
    f.name = FONT
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color
    return p


def agregar_parrafo_dual(tf, lead, resto, size=13.5, color=TEXT_COLOR,
                         space_after=6, bullet="\u2022  ", first=False):
    """Parrafo con encabezado en negrita (lead) y cuerpo normal (resto)."""
    p = tf.paragraphs[0] if first and not tf.paragraphs[0].runs else tf.add_paragraph()
    p.space_after = Pt(space_after)
    p.line_spacing = 1.04
    if bullet:
        run_b = p.add_run()
        run_b.text = bullet
        run_b.font.name = FONT
        run_b.font.size = Pt(size)
        run_b.font.color.rgb = color
    if lead:
        run_l = p.add_run()
        run_l.text = lead
        run_l.font.name = FONT
        run_l.font.size = Pt(size)
        run_l.font.bold = True
        run_l.font.color.rgb = color
    run_r = p.add_run()
    run_r.text = resto
    run_r.font.name = FONT
    run_r.font.size = Pt(size)
    run_r.font.color.rgb = color
    return p


# ---------------------------------------------------------------------------
# Utilidades de forma
# ---------------------------------------------------------------------------
def add_rect(slide, x, y, w, h, fill, line_color=None):
    """Rectangulo solido sin borde visible."""
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line_color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line_color
    shape.shadow.inherit = False
    return shape


def add_textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    """Caja de texto con word_wrap activo y anclaje configurable."""
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    return tf


def banda_titulo(slide, texto, size=26):
    """Banda azul oscuro superior con el titulo en blanco (bold)."""
    add_rect(slide, 0, 0, SLIDE_W, Inches(1.02), DARK_BLUE)
    box = slide.shapes.add_textbox(Inches(0.55), Inches(0.08),
                                   SLIDE_W - Inches(1.1), Inches(0.86))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = texto
    f = run.font
    f.name = FONT
    f.size = Pt(size)
    f.bold = True
    f.color.rgb = WHITE
    return box


def pie_diapositiva(slide, numero):
    """Pie con texto institucional y numero de diapositiva."""
    tf = add_textbox(slide, Inches(0.55), Inches(7.12), Inches(7.0), Inches(0.3))
    agregar_parrafo(tf, "PLN \u00b7 Entrega 2 \u00b7 Pontificia Universidad Javeriana",
                    size=9, color=LIGHT_GRAY, space_after=0, first=True)
    tf2 = add_textbox(slide, SLIDE_W - Inches(1.1), Inches(7.12), Inches(0.55), Inches(0.3))
    p = tf2.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    run = p.add_run()
    run.text = str(numero)
    run.font.name = FONT
    run.font.size = Pt(10)
    run.font.color.rgb = LIGHT_GRAY


def add_picture_scaled(slide, path, target_h_in, center_x_in, top_in):
    """Inserta un PNG con la altura pedida conservando su proporcion.

    Args:
        slide: Diapositiva destino.
        path: Ruta del PNG.
        target_h_in: Altura deseada en pulgadas.
        center_x_in: Centro horizontal en pulgadas.
        top_in: Borde superior en pulgadas.

    Returns:
        La forma picture creada.
    """
    w_px, h_px = png_size(path)
    w_in = target_h_in * (w_px / h_px)
    left = center_x_in - w_in / 2.0
    if not (0.0 <= left and left + w_in <= 13.333 and target_h_in > 0):
        raise ValueError(
            "Imagen fuera de rango: {} left={:.2f} right={:.2f} top={:.2f} h={:.2f}".format(
                path.name, left, left + w_in, top_in, target_h_in))
    return slide.shapes.add_picture(str(path), Inches(left), Inches(top_in),
                                    width=Inches(w_in), height=Inches(target_h_in))


def caption_figura(slide, texto, center_x_in, top_in, width_in=6.4):
    """Pie de figura pequeno en cursiva, centrado bajo la imagen."""
    tf = add_textbox(slide, Inches(center_x_in - width_in / 2.0), Inches(top_in),
                     Inches(width_in), Inches(0.3))
    agregar_parrafo(tf, texto, size=10.5, italic=True, color=GRAY_COLOR,
                    align=PP_ALIGN.CENTER, space_after=0, first=True)


def panel_con_encabezado(slide, x_in, y_in, w_in, h_in, titulo):
    """Panel con encabezado azul oscuro y fondo claro; devuelve su text frame."""
    add_rect(slide, Inches(x_in), Inches(y_in), Inches(w_in), Inches(h_in), LIGHT_BG)
    add_rect(slide, Inches(x_in), Inches(y_in), Inches(w_in), Inches(0.5), DARK_BLUE)
    tf_t = add_textbox(slide, Inches(x_in + 0.2), Inches(y_in + 0.03),
                       Inches(w_in - 0.4), Inches(0.44), anchor=MSO_ANCHOR.MIDDLE)
    agregar_parrafo(tf_t, titulo, size=14.5, bold=True, color=WHITE,
                    space_after=0, first=True)
    return add_textbox(slide, Inches(x_in + 0.25), Inches(y_in + 0.68),
                       Inches(w_in - 0.5), Inches(h_in - 0.85))


def set_cell(celda, texto, size, bold, color, fill, align=PP_ALIGN.LEFT):
    """Configura una celda de tabla con texto y estilo."""
    celda.fill.solid()
    celda.fill.fore_color.rgb = fill
    celda.margin_left = Inches(0.12)
    celda.margin_right = Inches(0.12)
    celda.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf_c = celda.text_frame
    p = tf_c.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = texto
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def notas_slide(slide, texto):
    """Establece las notas del orador de la diapositiva."""
    slide.notes_slide.notes_text_frame.text = texto


def verificar_ajuste(slide, indice, advertencias):
    """Advertencia (no bloqueante) si un text box parece no caber en su caja."""
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        tf = shape.text_frame
        texto = tf.text
        if not texto.strip():
            continue
        alto_in = shape.height / 914400.0
        ancho_in = shape.width / 914400.0
        alto_est = 0.0
        for p in tf.paragraphs:
            size_pt = max((r.font.size.pt if r.font.size else 14.0) for r in p.runs) \
                if p.runs else 14.0
            n_lineas = est_lineas(p.text, size_pt, ancho_in)
            spacing = p.line_spacing if p.line_spacing else 1.0
            alto_est += n_lineas * size_pt * 1.22 * spacing / 72.0
            alto_est += (p.space_after.pt if p.space_after else 0.0) / 72.0
        if alto_est > alto_in * 1.06:
            advertencias.append(
                "Slide {}: textbox probable desborde (est {} in > caja {} in): {}...".format(
                    indice, round(alto_est, 2), round(alto_in, 2),
                    texto[:60].replace("\n", " ")))


# ---------------------------------------------------------------------------
# Construccion de las 12 diapositivas
# ---------------------------------------------------------------------------
def construir(prs: Presentation) -> list[str]:
    """Construye todas las diapositivas; devuelve advertencias de layout."""
    advertencias: list[str] = []
    blank = prs.slide_layouts[6]

    # =====================================================================
    # Diapositiva 1 - Portada
    # =====================================================================
    slide = prs.slides.add_slide(blank)
    notas_slide(slide, "Presentamos la Entrega 2 del proyecto de Procesamiento de "
                       "Lenguaje Natural sobre jurisprudencia constitucional "
                       "colombiana. La entrega profundiza el analisis del corpus, "
                       "entrega el conjunto de datos medido, define la representacion "
                       "de los textos y propone la metodologia de modelado.")
    add_rect(slide, 0, 0, SLIDE_W, Inches(0.28), SECONDARY_BLUE)
    add_rect(slide, 0, Inches(6.78), SLIDE_W, Inches(0.72), DARK_BLUE)

    tf = add_textbox(slide, Inches(0.8), Inches(0.52), SLIDE_W - Inches(1.6), Inches(0.4))
    agregar_parrafo(tf, "PONTIFICIA UNIVERSIDAD JAVERIANA", size=14,
                    color=SECONDARY_BLUE, align=PP_ALIGN.CENTER,
                    space_after=0, first=True)
    tf = add_textbox(slide, Inches(0.8), Inches(0.95), SLIDE_W - Inches(1.6), Inches(0.4))
    agregar_parrafo(tf, "Procesamiento de Lenguaje Natural \u00b7 Entrega 2",
                    size=15, color=GRAY_COLOR, align=PP_ALIGN.CENTER,
                    space_after=0, first=True)
    add_rect(slide, Inches(5.17), Inches(1.58), Inches(3.0), Inches(0.045), DARK_BLUE)

    tf = add_textbox(slide, Inches(0.8), Inches(1.85), SLIDE_W - Inches(1.6), Inches(1.6))
    agregar_parrafo(tf, "EDA avanzado, representaci\u00f3n de textos\ny propuesta metodol\u00f3gica",
                    size=32, bold=True, color=DARK_BLUE, align=PP_ALIGN.CENTER,
                    space_after=6, first=True, line_spacing=1.05)
    agregar_parrafo(tf, "An\u00e1lisis computacional de la jurisprudencia de la "
                        "Corte Constitucional de Colombia (1992\u20132021)",
                    size=16, color=TEXT_COLOR, align=PP_ALIGN.CENTER, space_after=0)

    tf = add_textbox(slide, Inches(0.8), Inches(3.62), SLIDE_W - Inches(1.6), Inches(0.4))
    agregar_parrafo(tf, "23.742 sentencias \u00b7 1992\u20132021 \u00b7 corpus p\u00fablico "
                        "Hugging Face (CC-BY-4.0)",
                    size=14, italic=True, color=SECONDARY_BLUE,
                    align=PP_ALIGN.CENTER, space_after=0, first=True)

    tf = add_textbox(slide, Inches(0.8), Inches(4.18), SLIDE_W - Inches(1.6), Inches(2.3))
    agregar_parrafo(tf, "Integrantes", size=13, color=LIGHT_GRAY,
                    align=PP_ALIGN.CENTER, space_after=10, first=True)
    for nombre in ["Erika Valeria Rivera Jim\u00e9nez",
                   "Juli\u00e1n Guillermo Carrillo Meneses",
                   "Carlos Andr\u00e9s Garnica Salazar",
                   "Oscar Andr\u00e9s Prieto Moreno"]:
        agregar_parrafo(tf, nombre, size=17, color=TEXT_COLOR,
                        align=PP_ALIGN.CENTER, space_after=5)

    tf = add_textbox(slide, Inches(0.8), Inches(6.86), SLIDE_W - Inches(1.6), Inches(0.55))
    agregar_parrafo(tf, "Octubre de 2026", size=16, color=WHITE,
                    align=PP_ALIGN.CENTER, space_after=0, first=True)
    verificar_ajuste(slide, 1, advertencias)

    # =====================================================================
    # Diapositiva 2 - Problema y objetivo
    # =====================================================================
    slide = prs.slides.add_slide(blank)
    banda_titulo(slide, "Problema y objetivo")
    notas_slide(slide, "El problema es de PLN aplicado: extraer conocimiento de un "
                       "acervo enorme y sin metadatos. El objetivo del proyecto es un "
                       "sistema de apoyo que valide la providencia, determine su area "
                       "tematica, identifique los articulos de la Constitucion que la "
                       "sustentan y la situe en su contexto jurisprudencial. Esta "
                       "entrega cubre los requisitos 1 a 5 y 7: EDA avanzado, dataset "
                       "medido, representacion y propuesta metodologica.")

    tf = add_textbox(slide, Inches(0.75), Inches(1.35), Inches(6.1), Inches(5.5))
    agregar_parrafo(tf, "Problema", size=15, bold=True, color=DARK_BLUE,
                    space_after=8, first=True)
    items_problema = [
        ("Corpus extenso y formulaico. ", "\u224811.000 palabras por documento; "
         "estructura formal, f\u00f3rmulas de estilo y citas repetidas."),
        ("Sin anotaciones. ", "el tipo de providencia y el a\u00f1o deben inferirse "
         "del texto; no hay etiquetas tem\u00e1ticas."),
        ("Lectura manual inviable. ", "23.742 documentos y \u2248261 millones de "
         "palabras; la informaci\u00f3n relevante se dispersa entre plantillas."),
        ("Recursos escasos. ", "la literatura documenta escasez de recursos legales "
         "anotados en espa\u00f1ol para Am\u00e9rica Latina."),
    ]
    for i, (lead, resto) in enumerate(items_problema):
        agregar_parrafo_dual(tf, lead, resto, size=13.5, space_after=10, first=False)
    agregar_parrafo_dual(tf, "Aporte. ", "un corpus depurado, medido y documentado "
                         "es un aporte en s\u00ed mismo.", size=13.5,
                         space_after=0, bullet="\u25aa  ")

    tf = add_textbox(slide, Inches(7.25), Inches(1.35), Inches(5.35), Inches(3.0))
    agregar_parrafo(tf, "Objetivo del proyecto", size=15, bold=True, color=DARK_BLUE,
                    space_after=8, first=True)
    objetivos = [
        "Validar que la entrada sea una providencia.",
        "Determinar su \u00e1rea tem\u00e1tica constitucional.",
        "Identificar los art\u00edculos de la Constituci\u00f3n Pol\u00edtica que la sustentan.",
        "Situarla en su contexto jurisprudencial.",
    ]
    for obj in objetivos:
        agregar_parrafo(tf, obj, size=13.5, space_after=7, bullet="\u2013  ")

    tf = panel_con_encabezado(slide, 7.25, 4.3, 5.35, 2.62,
                              "Esta entrega (requisitos 1\u20135 y 7)")
    for i, item in enumerate([
            "Correcciones y modificaciones vs E1.",
            "EDA avanzado sobre el corpus completo.",
            "Conjunto de datos medido (parquet + muestra).",
            "Representaci\u00f3n de los textos (TF-IDF sobre lemas).",
            "Propuesta metodol\u00f3gica de modelado.",
            "Alcance: 1992\u20132021 (2021 parcial: 20 documentos)."]):
        agregar_parrafo(tf, item, size=11.5, space_after=4,
                        bullet="\u2022  ", first=(i == 0))
    pie_diapositiva(slide, 2)
    verificar_ajuste(slide, 2, advertencias)

    # =====================================================================
    # Diapositiva 3 - Corpus y correcciones vs E1
    # =====================================================================
    slide = prs.slides.add_slide(blank)
    banda_titulo(slide, "Corpus y correcciones respecto a la Entrega 1")
    notas_slide(slide, "El corpus se mantiene identico en cifras: 23.742 sentencias "
                       "T/C/SU, 1992-2021, con 8 exclusiones documentadas. Los cambios "
                       "frente a la Entrega 1 son extensiones verificadas: dataset "
                       "medido, EDA extendido, siete figuras nuevas, unidad de analisis "
                       "validada, analisis por periodos y la representacion con la "
                       "propuesta metodologica.")

    tf = panel_con_encabezado(slide, 0.75, 1.3, 4.55, 5.35, "Corpus de an\u00e1lisis")
    filas_corpus = [
        ("23.742", "sentencias depuradas"),
        ("T 17.099 \u00b7 C 6.246 \u00b7 SU 397", "tutela \u00b7 constitucionalidad \u00b7 unificaci\u00f3n"),
        ("1992\u20132021", "2021 parcial: 20 documentos"),
        ("8 exclusiones", "6 autos \u00b7 2 fragmentos HTML"),
        ("99,97 %", "extracci\u00f3n de id y a\u00f1o por regex"),
        ("Hugging Face (CC-BY-4.0)", "relator\u00eda de la Corte Constitucional"),
    ]
    for i, (valor, etiqueta) in enumerate(filas_corpus):
        agregar_parrafo(tf, valor, size=15.5, bold=True, color=DARK_BLUE,
                        space_after=1, first=(i == 0))
        agregar_parrafo(tf, etiqueta, size=11.5, color=GRAY_COLOR, space_after=9)

    tf = add_textbox(slide, Inches(5.75), Inches(1.32), Inches(6.85), Inches(4.6))
    agregar_parrafo(tf, "Modificaciones reales respecto a la E1", size=15, bold=True,
                    color=DARK_BLUE, space_after=8, first=True)
    cambios = [
        ("Dataset medido. ", "parquet de 23.742 filas \u00d7 12 columnas con texto, "
         "texto limpio y m\u00e9tricas por documento."),
        ("EDA extendido. ", "n-gramas, colocaciones, lemas/stems/POS, MTLD, "
         "granularidad y recursos l\u00e9xicos."),
        ("Visualizaciones. ", "siete figuras nuevas; las de la E1 se conservan."),
        ("Unidad validada. ", "sentencia 97,17 % vs subsentencia 4,64 % "
         "(criterio pre-registrado de +15 pp)."),
        ("Periodos. ", "contraste l\u00e9xico 1992\u20132001 / 2002\u20132011 / "
         "2012\u20132021 ejecutado (log-odds)."),
        ("Representaci\u00f3n y metodolog\u00eda. ", "TF-IDF sobre lemas y propuesta "
         "de modelado definidas."),
    ]
    for lead, resto in cambios:
        agregar_parrafo_dual(tf, lead, resto, size=12.5, space_after=8)

    add_rect(slide, Inches(5.75), Inches(5.85), Inches(6.85), Inches(0.8), NOTE_BG)
    tf = add_textbox(slide, Inches(6.0), Inches(5.98), Inches(6.4), Inches(0.6))
    agregar_parrafo(tf, "Las cifras can\u00f3nicas de la E1 (23.742 documentos; "
                        "260.988.316 / 250.012.101 palabras) no requirieron "
                        "correcci\u00f3n.",
                    size=12, color=TEXT_COLOR, space_after=0, first=True)
    pie_diapositiva(slide, 3)
    verificar_ajuste(slide, 3, advertencias)

    # =====================================================================
    # Diapositiva 4 - Limpieza del texto (pipeline)
    # =====================================================================
    slide = prs.slides.add_slide(blank)
    banda_titulo(slide, "Limpieza del texto: pipeline reproducible")
    notas_slide(slide, "La limpieza tiene ocho pasos medibles y documentados, "
                       "aplicados al corpus completo. La misma transformacion se "
                       "aplica al corpus y a la lista de stopwords: quedan 306 formas "
                       "efectivas. El filtro de ruido descarta restos de claves de "
                       "relatoria y citas de pagina antes de contar.")

    tf = add_textbox(slide, Inches(0.75), Inches(1.25), Inches(11.8), Inches(0.7))
    agregar_parrafo_dual(tf, "Decisi\u00f3n. ", "ocho pasos medibles sobre el corpus "
                         "completo; la misma transformaci\u00f3n se aplica al corpus y "
                         "a las stopwords (306 formas efectivas).",
                         size=13.5, space_after=0, bullet="", first=True)

    pasos_izq = [
        ("1", "Conversi\u00f3n del texto a min\u00fasculas."),
        ("2", "Protecci\u00f3n temporal de la \u00f1 para que no se pierda en NFD."),
        ("3", "Normalizaci\u00f3n Unicode (NFD) y eliminaci\u00f3n de diacr\u00edticos "
              "(art\u00edculo \u2192 articulo)."),
        ("4", "Filtro de caracteres: solo letras [a-z\u00f1]; fuera d\u00edgitos, "
              "puntuaci\u00f3n y s\u00edmbolos."),
    ]
    pasos_der = [
        ("5", "Colapso de espacios m\u00faltiples."),
        ("6", "Tokenizaci\u00f3n por separaci\u00f3n en espacios (exacta tras limpiar)."),
        ("7", "Filtro de ruido: descarta t, c, p (claves y citas de p\u00e1gina) y "
              "la abreviatura art."),
        ("8", "Stopwords NLTK del espa\u00f1ol normalizadas con el mismo criterio: "
              "306 formas (84 acentuadas corregidas)."),
    ]
    for col_x, pasos in ((0.95, pasos_izq), (6.95, pasos_der)):
        tf = add_textbox(slide, Inches(col_x), Inches(2.2), Inches(5.5), Inches(3.4))
        for i, (num, texto) in enumerate(pasos):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(12)
            p.line_spacing = 1.05
            rn = p.add_run()
            rn.text = num + ".  "
            rn.font.name = FONT
            rn.font.size = Pt(15)
            rn.font.bold = True
            rn.font.color.rgb = SECONDARY_BLUE
            rt = p.add_run()
            rt.text = texto
            rt.font.name = FONT
            rt.font.size = Pt(13.5)
            rt.font.color.rgb = TEXT_COLOR

    add_rect(slide, Inches(0.75), Inches(5.9), Inches(11.8), Inches(0.95), NOTE_BG)
    tf = add_textbox(slide, Inches(1.05), Inches(6.05), Inches(11.2), Inches(0.7))
    agregar_parrafo(tf, "Por qu\u00e9 as\u00ed: la literatura muestra que el "
                        "preprocesamiento altera los resultados y depende del corpus y "
                        "la tarea (HaCohen-Kerner et al., 2020); cada paso queda "
                        "documentado y es reproducible.",
                    size=12.5, color=TEXT_COLOR, space_after=0, first=True)
    pie_diapositiva(slide, 4)
    verificar_ajuste(slide, 4, advertencias)

    # =====================================================================
    # Diapositiva 5 - Medidas globales y distribucion temporal
    # =====================================================================
    slide = prs.slides.add_slide(blank)
    banda_titulo(slide, "Medidas globales y distribuci\u00f3n temporal")
    notas_slide(slide, "Los conteos son reales sobre el corpus completo: 260 millones "
                       "de palabras crudas y 250 millones limpias. Las stopwords pesan "
                       "48,48 por ciento de los tokens limpios. El volumen anual es "
                       "estable entre 1999 y 2014 y desciende despues; 2021 es parcial. "
                       "Sin stopwords emerge el lexico juridico dominante.")

    tf = add_textbox(slide, Inches(0.75), Inches(1.3), Inches(5.5), Inches(5.6))
    agregar_parrafo(tf, "Medidas globales", size=15, bold=True, color=DARK_BLUE,
                    space_after=8, first=True)
    medidas = [
        ("Documentos. ", "23.742 (T 17.099 \u00b7 C 6.246 \u00b7 SU 397)."),
        ("Palabras crudas. ", "260.988.316."),
        ("Palabras limpias. ", "250.012.101."),
        ("Stopwords. ", "121.205.437 (48,48 % de las limpias)."),
        ("Sin stopwords. ", "128.806.664."),
        ("Promedios por documento. ", "10.992,68 crudas / 10.530,37 limpias."),
        ("Mediana de lemas. ", "7.116 (p25: 3.959; p75: 12.654); palabras limpias: "
         "mediana 7.158."),
        ("Vocabulario. ", "224.468 formas \u00b7 195.959 lemas \u00b7 121.696 stems."),
    ]
    for lead, resto in medidas:
        agregar_parrafo_dual(tf, lead, resto, size=12.5, space_after=6)
    agregar_parrafo(tf, "Distribuci\u00f3n temporal: estable 1999\u20132014 "
                        "(rango de 941 en 2013 a 992 en 2002); descenso 2015\u20132020; "
                        "2021 parcial (20 documentos).",
                    size=12, italic=True, color=GRAY_COLOR, space_after=0,
                    space_before=4)

    add_picture_scaled(slide, FIG_POR_ANIO, 2.2, 9.9, 1.35)
    caption_figura(slide, "Sentencias por a\u00f1o de emisi\u00f3n (1992\u20132021).",
                   9.9, 3.6, 6.0)
    add_picture_scaled(slide, FIG_TOP20_SIN, 3.0, 9.9, 3.95)
    caption_figura(slide, "Top-20 de palabras sin stopwords.",
                   9.9, 7.0, 4.0)
    pie_diapositiva(slide, 5)
    verificar_ajuste(slide, 5, advertencias)

    # =====================================================================
    # Diapositiva 6 - N-gramas y colocaciones
    # =====================================================================
    slide = prs.slides.add_slide(blank)
    banda_titulo(slide, "N-gramas y colocaciones")
    notas_slide(slide, "El corpus es fuertemente formulaico: accion de tutela es el "
                       "trigrama mas frecuente y el 82,7 por ciento de los trigramas "
                       "tiene palabra intermedia vacia. La validacion manual de 20 "
                       "colocaciones da 75 por ciento de formulas validas con la razon "
                       "de verosimilitud, frente a 55 por ciento de la frecuencia y 0 "
                       "por ciento del PMI; por eso la LLR sera la medida principal.")

    tf = add_textbox(slide, Inches(0.75), Inches(1.3), Inches(5.4), Inches(5.6))
    agregar_parrafo(tf, "N-gramas m\u00e1s frecuentes", size=15, bold=True,
                    color=DARK_BLUE, space_after=8, first=True)
    ngramas = [
        ("Bigramas. ", "corte constitucional (286.786) \u00b7 derechos fundamentales "
         "(241.467) \u00b7 seguridad social (131.119)."),
        ("Trigramas. ", "accion de tutela (344.893) \u00b7 articulo del decreto "
         "(89.352) \u00b7 juez de tutela (42.253)."),
        ("F\u00f3rmulas. ", "el 82,7 % de los trigramas guardados tiene palabra "
         "intermedia vac\u00eda (accion de tutela)."),
    ]
    for lead, resto in ngramas:
        agregar_parrafo_dual(tf, lead, resto, size=12.5, space_after=8)
    agregar_parrafo(tf, "Validaci\u00f3n manual (20 colocaciones)", size=15, bold=True,
                    color=DARK_BLUE, space_after=8, space_before=6)
    validacion = [
        ("Frecuencia. ", "55 % de f\u00f3rmulas v\u00e1lidas."),
        ("PMI. ", "0 %: ordena ruido y fragmentos no espa\u00f1oles."),
        ("LLR. ", "75 %: equilibra frecuencia y asociaci\u00f3n \u2192 medida principal."),
    ]
    for lead, resto in validacion:
        agregar_parrafo_dual(tf, lead, resto, size=12.5, space_after=6)

    add_picture_scaled(slide, FIG_NGRAMAS, 2.45, 9.9, 1.35)
    caption_figura(slide, "Top-15 de bigramas y trigramas.", 9.9, 3.85, 6.0)
    add_picture_scaled(slide, FIG_COLOCACIONES, 2.45, 9.9, 4.2)
    caption_figura(slide, "Colocaciones por raz\u00f3n de verosimilitud (top-10).",
                   9.9, 6.7, 6.0)
    pie_diapositiva(slide, 6)
    verificar_ajuste(slide, 6, advertencias)

    # =====================================================================
    # Diapositiva 7 - Diversidad lexica y morfologia
    # =====================================================================
    slide = prs.slides.add_slide(blank)
    banda_titulo(slide, "Diversidad l\u00e9xica (MTLD) y morfolog\u00eda")
    notas_slide(slide, "La diversidad lexica desciende de forma sostenida en las tres "
                       "decadas: de 64,32 a 58,53, con una pendiente de menos 0,32 por "
                       "ano. La lematizacion reduce el vocabulario de 224.468 formas a "
                       "195.959 lemas conservando la legibilidad, y la distribucion POS "
                       "es coherente con el registro juridico.")

    tf = add_textbox(slide, Inches(0.75), Inches(1.3), Inches(5.5), Inches(5.6))
    agregar_parrafo(tf, "Diversidad l\u00e9xica (MTLD)", size=15, bold=True,
                    color=DARK_BLUE, space_after=8, first=True)
    mtld_items = [
        ("Global. ", "61,03; por tipo: T 61,63 > C 59,54 > SU 58,88."),
        ("Periodos. ", "64,32 \u2192 60,37 \u2192 58,53 (1992\u20132001 / "
         "2002\u20132011 / 2012\u20132021)."),
        ("Tendencia. ", "pendiente \u22120,32 por a\u00f1o (r = \u22120,328; "
         "p < 0,001)."),
    ]
    for lead, resto in mtld_items:
        agregar_parrafo_dual(tf, lead, resto, size=12.5, space_after=7)
    agregar_parrafo(tf, "Morfolog\u00eda", size=15, bold=True, color=DARK_BLUE,
                    space_after=8, space_before=6)
    morf_items = [
        ("Vocabulario. ", "224.468 formas \u2192 195.959 lemas (87,3 %) \u2192 "
         "121.696 stems (54,2 %)."),
        ("POS. ", "NOUN 22,58 % \u00b7 ADP 20,45 % \u00b7 DET 15,57 %."),
        ("Lectura. ", "lematizar reduce el vocabulario manteniendo la legibilidad: "
         "base del TF-IDF."),
    ]
    for lead, resto in morf_items:
        agregar_parrafo_dual(tf, lead, resto, size=12.5, space_after=7)

    add_picture_scaled(slide, FIG_MTLD, 2.1, 9.9, 1.35)
    caption_figura(slide, "MTLD por tipo de providencia y por a\u00f1o.",
                   9.9, 3.5, 6.2)
    add_picture_scaled(slide, FIG_LEMMA_STEM, 2.1, 9.9, 3.85)
    caption_figura(slide, "Compresi\u00f3n del vocabulario: formas, lemas y stems.",
                   9.9, 6.0, 6.2)
    pie_diapositiva(slide, 7)
    verificar_ajuste(slide, 7, advertencias)

    # =====================================================================
    # Diapositiva 8 - Recursos lexicos y ontologias
    # =====================================================================
    slide = prs.slides.add_slide(blank)
    banda_titulo(slide, "Recursos l\u00e9xicos y ontolog\u00edas")
    notas_slide(slide, "Evaluamos cuatro candidatas internacionales. Ninguna cubre el "
                       "espanol juridico colombiano: EuroVoc es parcial, LKIF es un "
                       "patron de modelado, ELI/ECLI no son vocabularios tematicos y "
                       "LegalRuleML exige anotacion manual costosa. La decision es "
                       "derivar una taxonomia ligera de la propia Constitucion Politica.")

    tf = add_textbox(slide, Inches(0.75), Inches(1.3), Inches(6.5), Inches(5.6))
    agregar_parrafo(tf, "Candidatas evaluadas", size=15, bold=True, color=DARK_BLUE,
                    space_after=8, first=True)
    candidatas = [
        ("EuroVoc / MultiEURLEX. ", "tesauro multiling\u00fce (espa\u00f1ol) de derecho "
         "de la UE; multietiqueta; aplicabilidad parcial (no mapea la CP ni la tutela)."),
        ("LKIF. ", "ontolog\u00eda formal OWL de conceptos jur\u00eddicos; aplicabilidad "
         "baja; \u00fatil como patr\u00f3n de modelado."),
        ("ELI / ECLI. ", "esquemas de identificaci\u00f3n y metadatos; no son "
         "vocabularios tem\u00e1ticos (aplicabilidad nula como contenido)."),
        ("LegalRuleML / Akoma Ntoso. ", "marcado de reglas y estructura; baja a corto "
         "plazo (anotaci\u00f3n manual costosa)."),
    ]
    for lead, resto in candidatas:
        agregar_parrafo_dual(tf, lead, resto, size=12.5, space_after=10)

    tf = panel_con_encabezado(slide, 7.55, 1.3, 5.05, 4.6, "Decisi\u00f3n")
    for i, item in enumerate([
            "No se adopta ninguna ontolog\u00eda externa tal cual.",
            "Taxonom\u00eda ligera derivada de la Constituci\u00f3n Pol\u00edtica: "
            "t\u00edtulos y cap\u00edtulos \u2192 bloques tem\u00e1ticos.",
            "Art\u00edculos 1\u2013380 \u2192 etiquetas finas.",
            "Respaldo \u201cderecho a X\u201d para las tutelas.",
            "EuroVoc en espa\u00f1ol: solo referencia de comparaci\u00f3n."]):
        agregar_parrafo(tf, item, size=12.5, space_after=7,
                        bullet="\u2022  ", first=(i == 0))

    add_rect(slide, Inches(7.55), Inches(6.05), Inches(5.05), Inches(0.62), NOTE_BG)
    tf = add_textbox(slide, Inches(7.75), Inches(6.14), Inches(4.65), Inches(0.5))
    agregar_parrafo(tf, "Condiciona el dise\u00f1o de etiquetas de la propuesta "
                        "metodol\u00f3gica (diapositiva 11).",
                    size=11.5, color=TEXT_COLOR, space_after=0, first=True)
    pie_diapositiva(slide, 8)
    verificar_ajuste(slide, 8, advertencias)

    # =====================================================================
    # Diapositiva 9 - Dataset entregado
    # =====================================================================
    slide = prs.slides.add_slide(blank)
    banda_titulo(slide, "Conjunto de datos medido (requisito 3)")
    notas_slide(slide, "El dataset medido es un parquet de 23.742 filas por 12 "
                       "columnas con el texto, el texto limpio y las metricas por "
                       "documento. La construccion se valido con aserciones duras: sin "
                       "nulos, ids unicos y sumas reconciliadas con las cifras "
                       "canonicas. Se adjuntan muestra estratificada, diccionario de "
                       "datos y copia de lemas por documento.")

    tf = add_textbox(slide, Inches(0.75), Inches(1.3), Inches(6.6), Inches(5.5))
    agregar_parrafo(tf, "corpus_medido.parquet \u2014 23.742 filas \u00d7 12 columnas",
                    size=14.5, bold=True, color=DARK_BLUE, space_after=10, first=True)
    grupos = [
        ("Identificaci\u00f3n. ", "id \u00b7 tipo \u00b7 numero \u00b7 ano."),
        ("Textos. ", "texto (crudo) \u00b7 texto_limpio (normalizado; tokens de ruido "
         "filtrados)."),
        ("M\u00e9tricas. ", "n_chars_texto \u00b7 n_palabras_raw \u00b7 "
         "n_palabras_clean \u00b7 n_stopwords \u00b7 n_lemas \u00b7 mtld."),
        ("Adjuntos. ", "corpus_medido_muestra.csv (100 filas por tipo y d\u00e9cada) "
         "\u00b7 diccionario_datos.md (14 variables) \u00b7 lemas_por_doc.parquet."),
    ]
    for lead, resto in grupos:
        agregar_parrafo_dual(tf, lead, resto, size=13, space_after=10)

    tf = panel_con_encabezado(slide, 7.6, 1.3, 5.0, 5.0, "Verificaci\u00f3n")
    for i, item in enumerate([
            "23.742 filas \u00b7 23.742 ids \u00fanicos.",
            "Sin nulos en claves ni en n_lemas/mtld.",
            "Sumas reconciliadas:",
            "crudas 260.988.316 \u00b7 limpias 250.012.101",
            "stopwords 121.205.437 \u00b7 lemas 248.533.344",
            "caracteres 1.640.749.780",
            "Tama\u00f1o del archivo: 1.544,7 MB."]):
        agregar_parrafo(tf, item, size=12, space_after=6,
                        bullet="\u2022  " if not item.startswith("crudas") and
                        not item.startswith("stopwords") and
                        not item.startswith("caracteres") else "",
                        first=(i == 0))

    add_rect(slide, Inches(0.75), Inches(6.15), Inches(11.85), Inches(0.55), NOTE_BG)
    tf = add_textbox(slide, Inches(1.0), Inches(6.23), Inches(11.4), Inches(0.45))
    agregar_parrafo(tf, "El archivo se adjunta en la carpeta datos/ junto al informe "
                        "(formato legible y variables relevantes).",
                    size=11.5, color=TEXT_COLOR, space_after=0, first=True)
    pie_diapositiva(slide, 9)
    verificar_ajuste(slide, 9, advertencias)

    # =====================================================================
    # Diapositiva 10 - Representacion: TF-IDF y embeddings
    # =====================================================================
    slide = prs.slides.add_slide(blank)
    banda_titulo(slide, "Representaci\u00f3n: TF-IDF y embeddings")
    notas_slide(slide, "La representacion principal es TF-IDF sobre lemas, "
                       "interpretable y de bajo costo: la matriz tiene 95,66 por ciento "
                       "de esparsidad y se ajusto en poco mas de tres minutos. Los "
                       "embeddings contextuales se evaluaron como comparacion: capturan "
                       "cercania semantica con mayor costo y menor interpretabilidad.")

    tf = add_textbox(slide, Inches(0.75), Inches(1.3), Inches(5.6), Inches(5.6))
    agregar_parrafo(tf, "TF-IDF sobre lemas (representaci\u00f3n principal)",
                    size=14.5, bold=True, color=DARK_BLUE, space_after=8, first=True)
    tfidf_items = [
        ("Configuraci\u00f3n. ", "1\u20132 gramos \u00b7 min_df=3 \u00b7 100.000 "
         "rasgos \u00b7 frecuencia sublineal."),
        ("Matriz. ", "23.742 \u00d7 100.000 \u00b7 nnz 103.045.759 \u00b7 esparsidad "
         "95,66 % \u00b7 ajuste 192,5 s."),
        ("Cola larga. ", "\u22482 % de la masa de frecuencia documental en los 100 "
         "primeros rasgos, \u224817 % en los 1.000 y \u224860 % en los 10.000."),
        ("Ejemplo T-001/2021. ", "7.712 rasgos activos; mayor peso: capital salud "
         "(0,065) \u00b7 valero (0,056) \u00b7 de calle (0,052)."),
        ("Similitud T/C/SU. ", "0,24\u20130,25 con TF-IDF: discrimina por "
         "solapamiento l\u00e9xico."),
    ]
    for lead, resto in tfidf_items:
        agregar_parrafo_dual(tf, lead, resto, size=12.5, space_after=8)
    agregar_parrafo(tf, "Embeddings contextuales (comparaci\u00f3n)", size=14.5,
                    bold=True, color=DARK_BLUE, space_after=8, space_before=6)
    emb_items = [
        ("e5 multiling\u00fce. ", "768 dimensiones \u00b7 5 documentos \u00b7 512 "
         "tokens \u00b7 GPU."),
        ("Similitud media. ", "0,899; en los mismos documentos 0,91\u20130,92 "
         "\u2192 complementarios (mayor costo, menor interpretabilidad)."),
    ]
    for lead, resto in emb_items:
        agregar_parrafo_dual(tf, lead, resto, size=12.5, space_after=8)

    add_picture_scaled(slide, FIG_TFIDF, 2.5, 9.9, 1.35)
    caption_figura(slide, "TF-IDF: cola larga del vocabulario y top-15 rasgos del "
                          "documento T-001/2021.", 9.9, 3.95, 6.2)
    add_picture_scaled(slide, FIG_EMBEDDINGS, 2.6, 9.9, 4.35)
    caption_figura(slide, "Similitud coseno e5-base (5 documentos).",
                   9.9, 7.05, 4.5)
    pie_diapositiva(slide, 10)
    verificar_ajuste(slide, 10, advertencias)

    # =====================================================================
    # Diapositiva 11 - Propuesta metodologica
    # =====================================================================
    slide = prs.slides.add_slide(blank)
    banda_titulo(slide, "Propuesta metodol\u00f3gica de los modelos")
    notas_slide(slide, "La tarea es una clasificacion supervisada multi-clase: el tipo "
                       "de providencia funciona como categoria inicial y el area "
                       "tematica derivada de la Constitucion como objetivo principal, "
                       "con etiquetas por validar. Las particiones son 70/15/15 "
                       "estratificadas con prueba intocable y criterios anti-leakage. "
                       "Se compararan la linea base mayoritaria, TF-IDF con modelos "
                       "lineales y un transformer en espanol, con macro-F1 y analisis "
                       "de errores como criterios.")

    tf = add_textbox(slide, Inches(0.75), Inches(1.28), Inches(11.8), Inches(1.2))
    agregar_parrafo_dual(tf, "Tarea. ", "clasificaci\u00f3n supervisada multi-clase: "
                         "tipo de providencia (T/C/SU) como categor\u00eda inicial; "
                         "\u00e1rea tem\u00e1tica derivada de la CP como objetivo "
                         "principal (etiquetas a validar).",
                         size=12.5, space_after=2, bullet="", first=True)
    agregar_parrafo_dual(tf, "Particiones. ", "70/15/15 estratificado (semilla 42, "
                         "prueba intocable) \u00b7 vectorizador ajustado solo con "
                         "entrenamiento \u00b7 partici\u00f3n a nivel de documento \u00b7 "
                         "evaluaci\u00f3n cronol\u00f3gica complementaria.",
                         size=12.5, space_after=0, bullet="", first=False)

    shape_tabla = slide.shapes.add_table(4, 3, Inches(0.75), Inches(2.75),
                                         Inches(11.8), Inches(2.5))
    tabla = shape_tabla.table
    tabla.columns[0].width = Inches(2.9)
    tabla.columns[1].width = Inches(3.2)
    tabla.columns[2].width = Inches(5.7)
    tabla.rows[0].height = Inches(0.5)
    for i in range(1, 4):
        tabla.rows[i].height = Inches(0.62)
    set_cell(tabla.cell(0, 0), "Enfoque", 13, True, WHITE, DARK_BLUE)
    set_cell(tabla.cell(0, 1), "Representaci\u00f3n", 13, True, WHITE, DARK_BLUE)
    set_cell(tabla.cell(0, 2), "Fortalezas y criterios", 13, True, WHITE, DARK_BLUE)
    filas_modelos = [
        ("L\u00ednea base mayoritaria", "\u2014",
         "Piso de comparaci\u00f3n; macro-F1 y F1 por clase."),
        ("Cl\u00e1sica: modelos lineales", "TF-IDF sobre lemas (1\u20132 gramos)",
         "Interpretable, bajo costo; macro-F1, matriz de confusi\u00f3n, an\u00e1lisis "
         "de errores, costo."),
        ("Transformer en espa\u00f1ol", "BETO con segmentaci\u00f3n y agregaci\u00f3n",
         "Contexto y composicionalidad; costo de GPU y ventana de 512; robustez por "
         "periodo."),
    ]
    for i, (enf, rep, crit) in enumerate(filas_modelos, start=1):
        fill = LIGHT_BG if i % 2 == 1 else WHITE
        set_cell(tabla.cell(i, 0), enf, 12, True, DARK_BLUE, fill)
        set_cell(tabla.cell(i, 1), rep, 12, False, TEXT_COLOR, fill)
        set_cell(tabla.cell(i, 2), crit, 11.5, False, TEXT_COLOR, fill)

    add_rect(slide, Inches(0.75), Inches(5.55), Inches(11.8), Inches(1.15), NOTE_BG)
    tf = add_textbox(slide, Inches(1.05), Inches(5.72), Inches(11.2), Inches(0.85))
    agregar_parrafo(tf, "Criterios de evaluaci\u00f3n: macro-F1 (principal, por el "
                        "desbalance T 72,02 %) \u00b7 F1 por clase \u00b7 matriz de "
                        "confusi\u00f3n \u00b7 an\u00e1lisis cualitativo de errores \u00b7 "
                        "robustez por periodo y tipo \u00b7 costo (tiempo, memoria, GPU).",
                    size=12.5, color=TEXT_COLOR, space_after=0, first=True)
    pie_diapositiva(slide, 11)
    verificar_ajuste(slide, 11, advertencias)

    # =====================================================================
    # Diapositiva 12 - Conclusiones y trabajo futuro
    # =====================================================================
    slide = prs.slides.add_slide(blank)
    banda_titulo(slide, "Conclusiones y trabajo futuro")
    notas_slide(slide, "Cerramos con los cinco hallazgos principales: corpus y dataset "
                       "consolidados, genero formulaico con LLR como mejor medida, "
                       "diversidad decreciente con recambio tematico, unidad de "
                       "analisis validada y representacion con metodologia definidas. "
                       "El plan inmediato: validar etiquetas, comparar modelos, "
                       "analizar errores y explorar embeddings y resumenes.")

    tf = add_textbox(slide, Inches(0.75), Inches(1.3), Inches(6.0), Inches(5.4))
    agregar_parrafo(tf, "Conclusiones", size=15, bold=True, color=DARK_BLUE,
                    space_after=8, first=True)
    conclusiones = [
        ("Corpus y dataset. ", "23.742 sentencias consolidadas; dataset medido "
         "verificado (12 columnas); cifras E1 sin correcci\u00f3n."),
        ("G\u00e9nero formulaico. ", "accion de tutela (344.893) domina; LLR 75 % como "
         "mejor medida de colocaciones."),
        ("Diversidad y diacron\u00eda. ", "MTLD 64,32 \u2192 58,53; recambio tem\u00e1tico: "
         "incoder/rupd \u2192 uariv/jep/coronavirus."),
        ("Unidad validada. ", "sentencia 97,17 % vs subsentencia 4,64 %."),
        ("Representaci\u00f3n y m\u00e9todo. ", "TF-IDF sobre lemas calculado (95,66 % de "
         "esparsidad); propuesta metodol\u00f3gica definida."),
    ]
    for lead, resto in conclusiones:
        agregar_parrafo_dual(tf, lead, resto, size=12.5, space_after=10)

    tf = add_textbox(slide, Inches(7.15), Inches(1.3), Inches(5.45), Inches(4.6))
    agregar_parrafo(tf, "Trabajo futuro", size=15, bold=True, color=DARK_BLUE,
                    space_after=8, first=True)
    futuro = [
        "Validar las etiquetas tem\u00e1ticas con personas expertas.",
        "Comparar modelos (lineales vs transformer) con macro-F1 y an\u00e1lisis de errores.",
        "Evaluar embeddings de dominio y recuperaci\u00f3n de fragmentos.",
        "Explorar res\u00famenes y extracci\u00f3n para la consulta de sentencias extensas.",
    ]
    for item in futuro:
        agregar_parrafo(tf, item, size=12.5, space_after=9, bullet="\u2022  ")

    add_rect(slide, Inches(7.15), Inches(5.35), Inches(5.45), Inches(1.35), NOTE_BG)
    tf = add_textbox(slide, Inches(7.4), Inches(5.5), Inches(4.95), Inches(1.05))
    agregar_parrafo(tf, "El dataset medido, el EDA avanzado y la metodolog\u00eda "
                        "dejan al proyecto listo para su fase de modelado.",
                    size=12, italic=True, color=SECONDARY_BLUE, space_after=0,
                    first=True)

    tf = add_textbox(slide, Inches(0.75), Inches(6.3), Inches(6.0), Inches(0.6))
    agregar_parrafo(tf, "Gracias \u00b7 Preguntas", size=13, bold=True,
                    color=SECONDARY_BLUE, space_after=0, first=True)
    pie_diapositiva(slide, 12)
    verificar_ajuste(slide, 12, advertencias)

    return advertencias


# ---------------------------------------------------------------------------
# Verificacion del archivo generado
# ---------------------------------------------------------------------------
def verificar_pptx(path: Path) -> dict:
    """Reabre el .pptx y verifica diapositivas, imagenes y notas.

    Args:
        path: Ruta del archivo generado.

    Returns:
        Diccionario con los conteos verificados.
    """
    prs = Presentation(str(path))
    n_slides = len(prs.slides._sldIdLst)
    n_pics = 0
    n_notas = 0
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                n_pics += 1
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            n_notas += 1
    return {"diapositivas": n_slides, "imagenes": n_pics, "notas": n_notas}


def main() -> None:
    """Genera el pptx, lo guarda y verifica el resultado."""
    print("Verificando figuras de entrada...")
    for ruta in FIGURAS_USADAS:
        esperar_figura(ruta)
        print("  OK {}".format(ruta.name))

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    prs.core_properties.title = ("Entrega 2 - EDA avanzado, representacion de textos "
                                 "y propuesta metodologica")
    prs.core_properties.subject = "Procesamiento de Lenguaje Natural - PUJ"
    prs.core_properties.author = "Rivera, Carrillo, Garnica y Prieto"

    advertencias = construir(prs)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT_FILE))

    print()
    print("Presentacion guardada en: {}".format(OUT_FILE))
    resumen = verificar_pptx(OUT_FILE)
    print("Diapositivas verificadas: {}".format(resumen["diapositivas"]))
    print("Imagenes embebidas verificadas: {}".format(resumen["imagenes"]))
    print("Diapositivas con notas del orador: {}".format(resumen["notas"]))
    print("Tamano: {:.2f} MB".format(OUT_FILE.stat().st_size / 1e6))
    if advertencias:
        print("ADVERTENCIAS DE LAYOUT (revisar):")
        for a in advertencias:
            print("  - " + a)
    else:
        print("Sin advertencias de desborde de texto estimadas.")


if __name__ == "__main__":
    main()
